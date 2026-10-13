from __future__ import annotations

import csv
import html
import io
import json
import os
import re
from html.parser import HTMLParser
from pathlib import Path

from wemute.engines.common import ConversionError, Reporter, last_line, run_tool

Sheet = tuple[str, list[list[object]]]

PANDOC_TIMEOUT = 180
_NUMBER = re.compile(r"^-?(?:0|[1-9]\d{0,14})(?:\.\d+)?$")


# ---------------------------------------------------------------- reading


def decode_text(data: bytes) -> str:
    """UTF-8 first, then GB18030 (common for Chinese CSV exports), then Latin-1."""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_delimited(source: Path, suffix: str) -> list[list[object]]:
    text = decode_text(source.read_bytes())
    if suffix == ".tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_workbook(source: Path) -> list[Sheet]:
    import openpyxl

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        sheets: list[Sheet] = []
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            while rows and all(value in (None, "") for value in rows[-1]):
                rows.pop()
            width = max((len(row) for row in rows), default=0)
            # Trim trailing empty columns that spreadsheets often carry.
            while width and all(len(row) < width or row[width - 1] in (None, "") for row in rows):
                width -= 1
            sheets.append((sheet.title, [row[:width] for row in rows]))
        return sheets
    finally:
        workbook.close()


def _flatten(value: object, prefix: str = "") -> dict[str, object]:
    if isinstance(value, dict):
        flat: dict[str, object] = {}
        for key, item in value.items():
            name = f"{prefix}.{key}" if prefix else str(key)
            flat.update(_flatten(item, name))
        return flat
    if isinstance(value, list):
        return {prefix or "value": json.dumps(value, ensure_ascii=False)}
    return {prefix or "value": value}


def json_records(payload: object) -> list[dict[str, object]]:
    """Find the rows in a JSON document: a list, or the one list inside an object."""
    if isinstance(payload, dict):
        lists = [value for value in payload.values() if isinstance(value, list) and value]
        if len(lists) == 1 and all(isinstance(item, dict) for item in lists[0]):
            payload = lists[0]
        elif payload and all(isinstance(value, dict) for value in payload.values()):
            payload = [{"key": key, **value} for key, value in payload.items()]
        else:
            payload = [payload]
    if not isinstance(payload, list):
        return []
    records: list[dict[str, object]] = []
    for item in payload:
        if isinstance(item, dict):
            records.append(_flatten(item))
        elif isinstance(item, list):
            records.append({f"column {index + 1}": value for index, value in enumerate(item)})
        else:
            records.append({"value": item})
    return records


def read_json_table(source: Path, reporter: Reporter) -> list[Sheet]:
    try:
        payload = json.loads(decode_text(source.read_bytes()))
    except json.JSONDecodeError as exc:
        raise reporter.error("bad_json", detail=f"{exc.msg} (line {exc.lineno})") from exc
    records = json_records(payload)
    if not records:
        raise reporter.error("no_table")
    header: list[str] = []
    for record in records:
        for key in record:
            if key not in header:
                header.append(key)
    rows: list[list[object]] = [list(header)]
    rows.extend([record.get(key, "") for key in header] for record in records)
    return [(source.stem, rows)]


def read_table(source: Path, suffix: str, reporter: Reporter) -> list[Sheet]:
    if suffix in {".csv", ".tsv"}:
        return [(source.stem, read_delimited(source, suffix))]
    if suffix == ".json":
        return read_json_table(source, reporter)
    return read_workbook(source)


# ---------------------------------------------------------------- writing


def typed(value: object) -> object:
    """Numbers stay numbers; IDs with leading zeros or 16+ digits stay text."""
    if isinstance(value, str) and _NUMBER.match(value.strip()):
        stripped = value.strip()
        return float(stripped) if "." in stripped else int(stripped)
    return value


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
    return str(value)


def write_csv(rows: list[list[object]], output: Path) -> None:
    # UTF-8 with BOM so Excel and Numbers open Chinese text correctly.
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        for row in rows:
            writer.writerow([_cell_text(value) for value in row])


def write_xlsx(sheets: list[Sheet], output: Path) -> None:
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook(write_only=True)
    bold = Font(bold=True)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title=re.sub(r"[\[\]:*?/\\]", "_", title)[:31] or "Sheet")
        widths: dict[int, int] = {}
        for row in rows[:200]:
            for index, value in enumerate(row):
                widths[index] = max(widths.get(index, 0), min(48, len(_cell_text(value)) + 2))
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index + 1)].width = max(8, width)
        if rows:
            sheet.freeze_panes = "A2"
        for number, row in enumerate(rows):
            if number == 0:
                cells = []
                for value in row:
                    cell = WriteOnlyCell(sheet, value=_cell_text(value))
                    cell.font = bold
                    cells.append(cell)
                sheet.append(cells)
            else:
                sheet.append([typed(value) for value in row])
    workbook.save(output)


def _records(rows: list[list[object]]) -> list[dict[str, object]]:
    if not rows:
        return []
    header = [_cell_text(value) or f"column {index + 1}" for index, value in enumerate(rows[0])]
    records = []
    for row in rows[1:]:
        values = [typed(value) if isinstance(value, str) else value for value in row]
        values += [None] * (len(header) - len(values))
        records.append(
            {
                key: (_cell_text(value) if hasattr(value, "isoformat") else value)
                for key, value in zip(header, values)
            }
        )
    return records


def write_json(sheets: list[Sheet], output: Path) -> None:
    if len(sheets) == 1:
        payload: object = _records(sheets[0][1])
    else:
        payload = {title: _records(rows) for title, rows in sheets}
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _markdown_cell(value: object) -> str:
    return _cell_text(value).replace("|", "\\|").replace("\n", "<br>")


def rows_to_markdown(rows: list[list[object]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [list(row) + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(_markdown_cell(value) for value in padded[0]) + " |"]
    lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
    lines.extend("| " + " | ".join(_markdown_cell(value) for value in row) + " |" for row in padded[1:])
    return "\n".join(lines) + "\n"


def rows_to_html_table(rows: list[list[object]]) -> str:
    if not rows:
        return "<table></table>"
    head = "".join(f"<th>{html.escape(_cell_text(value))}</th>" for value in rows[0])
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(_cell_text(value))}</td>" for value in row) + "</tr>"
        for row in rows[1:]
    )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def sheet_file_name(stem: str, title: str, extension: str) -> str:
    safe = re.sub(r'[/:\\]', "-", title).strip() or "Sheet"
    return f"{stem} - {safe}{extension}"


def table_convert(
    source: Path,
    suffix: str,
    output: Path,
    target: str,
    reporter: Reporter,
    stem: str | None = None,
) -> None:
    """CSV/TSV/XLSX/JSON → CSV, XLSX, JSON, MD or HTML.

    A workbook with several sheets becomes a folder of CSVs at ``output``.
    """
    reporter.stage("stage_reading")
    sheets = read_table(source, suffix, reporter)
    reporter.progress(0.5)
    reporter.stage("stage_saving")
    stem = stem or source.stem
    if target == "CSV":
        filled = [(title, rows) for title, rows in sheets if rows] or sheets[:1]
        if len(filled) == 1:
            write_csv(filled[0][1], output)
        else:
            output.mkdir()
            for title, rows in filled:
                write_csv(rows, output / sheet_file_name(stem, title, ".csv"))
    elif target == "XLSX":
        write_xlsx(sheets, output)
    elif target == "JSON":
        write_json(sheets, output)
    elif target == "MD":
        parts = [
            (f"## {title}\n\n" if len(sheets) > 1 else "") + rows_to_markdown(rows)
            for title, rows in sheets
        ]
        output.write_text("\n".join(parts), encoding="utf-8")
    elif target == "HTML":
        tables = "".join(
            (f"<h2>{html.escape(title)}</h2>" if len(sheets) > 1 else "") + rows_to_html_table(rows)
            for title, rows in sheets
        )
        output.write_text(html_document(stem, tables), encoding="utf-8")
    else:
        raise ConversionError(f"Unsupported table output: {target}")
    reporter.progress(1.0)


# ------------------------------------------------------------------- text

_STYLE = """
body { font-family: -apple-system, "Helvetica Neue", "PingFang SC", sans-serif;
       line-height: 1.6; max-width: 46em; margin: 2em auto; padding: 0 1em; color: #182033; }
table { border-collapse: collapse; margin: 1em 0; }
th, td { border: 1px solid #d5dae6; padding: 4px 10px; text-align: left; }
th { background: #f2f5fb; }
pre, code { font-family: "SF Mono", Menlo, monospace; font-size: 0.92em; }
pre { background: #f4f6fb; padding: 0.8em 1em; border-radius: 6px; overflow-x: auto; }
img { max-width: 100%; }
"""


def html_document(title: str, body: str) -> str:
    return (
        "<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\">"
        f"<title>{html.escape(title)}</title><style>{_STYLE}</style></head>"
        f"<body>\n{body}\n</body></html>\n"
    )


def markdown_to_html_body(text: str) -> str:
    import markdown

    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists", "toc"])


class _TextExtractor(HTMLParser):
    _BLOCKS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
               "pre", "blockquote", "section", "article", "table", "ul", "ol", "hr"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0
        self._pre = 0
        self._cell = 0

    def _block_break(self) -> None:
        self.parts.append(" " if self._cell else "\n")

    def handle_starttag(self, tag: str, attrs) -> None:  # noqa: ANN001
        if tag in {"script", "style", "head", "title"}:
            self._skip += 1
        elif tag in self._BLOCKS:
            self._block_break()
        if tag == "pre":
            self._pre += 1
        if tag == "li":
            self.parts.append("• ")
        if tag in {"td", "th"}:
            self.parts.append("\t")
            self._cell += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "head", "title"}:
            self._skip = max(0, self._skip - 1)
        elif tag in {"td", "th"}:
            self._cell = max(0, self._cell - 1)
        elif tag in self._BLOCKS:
            self._block_break()
        if tag == "pre":
            self._pre = max(0, self._pre - 1)

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data if self._pre else re.sub(r"\s+", " ", data))

    def text(self) -> str:
        joined = "".join(self.parts)
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in joined.splitlines()]
        return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip() + "\n"


def html_to_text(markup: str) -> str:
    parser = _TextExtractor()
    parser.feed(markup)
    parser.close()
    return parser.text()


def text_to_html_body(text: str) -> str:
    paragraphs = [block.strip("\n") for block in re.split(r"\n\s*\n", text) if block.strip()]
    return "\n".join(
        "<p>" + html.escape(block).replace("\n", "<br>\n") + "</p>" for block in paragraphs
    )


def write_pdf_from_html(
    markup: str, output: Path, base: Path, reporter: Reporter, title: str = ""
) -> None:
    """Lay out HTML with Qt's rich-text engine on A4 pages (no network access)."""
    from PySide6.QtCore import QMarginsF, QUrl
    from PySide6.QtGui import QFont, QPageLayout, QPageSize, QPdfWriter, QTextDocument

    from wemute.engines.qt import ensure_gui

    ensure_gui()
    reporter.stage("stage_rendering")
    document = QTextDocument()
    document.setBaseUrl(QUrl.fromLocalFile(str(base) + "/"))
    document.setDefaultFont(QFont("PingFang SC", 11))
    document.setDefaultStyleSheet(
        "table { border-collapse: collapse; } th, td { border: 1px solid #c8cedb; padding: 4px; }"
        " pre { background: #f4f6fb; } h1, h2, h3 { color: #182033; }"
    )
    document.setHtml(markup)
    writer = QPdfWriter(str(output))
    writer.setPageLayout(
        QPageLayout(
            QPageSize(QPageSize.PageSizeId.A4),
            QPageLayout.Orientation.Portrait,
            QMarginsF(18, 18, 18, 18),
            QPageLayout.Unit.Millimeter,
        )
    )
    writer.setTitle(title or output.stem)
    document.print_(writer)
    del writer
    if not output.is_file() or output.stat().st_size == 0:
        raise reporter.error("no_pages")


def write_docx_from_text(text: str, output: Path) -> None:
    from docx import Document

    document = Document()
    for block in re.split(r"\n\s*\n", text):
        if block.strip():
            document.add_paragraph(block.strip("\n"))
    document.save(output)


def pandoc_convert(
    pandoc: Path,
    source: Path,
    output: Path,
    reader: str,
    writer: str,
    reporter: Reporter,
) -> None:
    if not pandoc.is_file():
        raise reporter.error("tool_missing", name="Pandoc")
    reporter.stage("stage_converting")
    reporter.indeterminate(True)
    command = [
        str(pandoc),
        "--from",
        reader,
        "--to",
        writer,
        "--resource-path",
        str(source.parent),
        "--output",
        str(output),
        str(source),
    ]
    if writer in {"gfm", "markdown"}:
        command.insert(1, "--wrap=none")
    result = run_tool(
        command,
        timeout=PANDOC_TIMEOUT,
        timeout_message=reporter.text("tool_timeout", seconds=PANDOC_TIMEOUT),
        cwd=source.parent,
        env={**os.environ, "LANG": "en_US.UTF-8"},
    )
    if result.returncode != 0 or not output.is_file():
        detail = last_line(result.stdout)
        message = reporter.text("pandoc_failed")
        raise ConversionError(f"{message}: {detail}" if detail else message)


def text_convert(
    source: Path,
    suffix: str,
    output: Path,
    target: str,
    reporter: Reporter,
    pandoc: Path | None = None,
) -> None:
    """Markdown, HTML and plain text to DOCX, HTML, MD, PDF or TXT."""
    reporter.stage("stage_reading")
    text = decode_text(source.read_bytes())
    kind = "md" if suffix in {".md", ".markdown"} else ("html" if suffix in {".html", ".htm"} else "txt")
    reporter.progress(0.2)
    if target == "DOCX" and kind != "txt":
        pandoc_convert(pandoc or Path(), source, output, "gfm" if kind == "md" else "html", "docx", reporter)
    elif target == "MD":
        pandoc_convert(pandoc or Path(), source, output, "html", "gfm", reporter)
    elif target == "DOCX":
        write_docx_from_text(text, output)
    elif target == "HTML":
        body = markdown_to_html_body(text) if kind == "md" else text_to_html_body(text)
        output.write_text(html_document(source.stem, body), encoding="utf-8")
    elif target == "TXT":
        markup = markdown_to_html_body(text) if kind == "md" else text
        output.write_text(html_to_text(markup), encoding="utf-8")
    elif target == "PDF":
        if kind == "md":
            markup = html_document(source.stem, markdown_to_html_body(text))
        elif kind == "html":
            markup = text
        else:
            markup = html_document(source.stem, text_to_html_body(text))
        write_pdf_from_html(markup, output, source.parent, reporter, source.stem)
    else:
        raise ConversionError(f"Unsupported text output: {target}")
    reporter.progress(1.0)
