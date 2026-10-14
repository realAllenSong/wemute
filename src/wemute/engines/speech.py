"""Subtitles, transcripts and read-aloud through ThunderTalk's headless CLI.

ThunderTalk (the owner's local voice app) does the speech work on this Mac;
Wemute runs ``ThunderTalk --cli transcribe|speak`` and turns its JSON lines
into the worker's progress protocol.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from wemute.engines.common import ConversionError, Reporter, remove_tree, run_tool, scratch_dir

# Long recordings take a while on the slower models; ThunderTalk reports
# progress throughout, so this only catches a stuck process.
TIMEOUT_SECONDS = 4 * 3600

_STAGES = {
    "load_model": "stage_loading_speech",
    "transcribe": "stage_transcribing",
    "speak": "stage_speaking",
    "write": "stage_saving",
}

# ThunderTalk's stable error codes → Wemute messages.
_ERRORS = {
    "no_speech": "speech_none",
    "no_model": "speech_no_model",
    "tts_missing": "speech_no_voice",
    "permission_denied": "speech_permission",
    "locale_unsupported": "speech_locale",
    "decode": "speech_decode",
    "ffmpeg_missing": "speech_decode",
}

_MARKDOWN = [
    (re.compile(r"```.*?```", re.S), " "),
    (re.compile(r"!\[[^\]]*\]\([^)]*\)"), " "),
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),
    (re.compile(r"^\s{0,3}#{1,6}\s*", re.M), ""),
    (re.compile(r"^\s*[-*+>]\s+", re.M), ""),
    (re.compile(r"[*_`~]{1,3}"), ""),
]


def text_for_reading(source: Path, suffix: str) -> str:
    """Plain text to read aloud from a text, Markdown, HTML, Word or PDF file."""
    suffix = suffix.lower()
    if suffix == ".pdf":
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(str(source))
        try:
            pages = [document[index].get_textpage().get_text_range() for index in range(len(document))]
        finally:
            document.close()
        return "\n\n".join(pages)
    if suffix == ".docx":
        import docx

        return "\n".join(paragraph.text for paragraph in docx.Document(str(source)).paragraphs)
    if suffix in {".txt", ".md", ".markdown", ".html", ".htm"}:
        text = source.read_text(encoding="utf-8", errors="replace")
        if suffix in {".html", ".htm"}:
            from html.parser import HTMLParser

            parts: list[str] = []

            class _Text(HTMLParser):
                skip = 0

                def handle_starttag(self, tag, attrs):  # noqa: ANN001, ARG002
                    if tag in {"script", "style"}:
                        self.skip += 1

                def handle_endtag(self, tag):  # noqa: ANN001
                    if tag in {"script", "style"} and self.skip:
                        self.skip -= 1
                    elif tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr"}:
                        parts.append("\n")

                def handle_data(self, data):  # noqa: ANN001
                    if not self.skip:
                        parts.append(data)

            _Text().feed(text)
            return "".join(parts)
        if suffix in {".md", ".markdown"}:
            for pattern, replacement in _MARKDOWN:
                text = pattern.sub(replacement, text)
        return text
    raise ValueError(f"cannot read text from {suffix} files")


# The owner's choice for reading aloud; the best ready voice when it isn't installed.
READ_ALOUD_ENGINE = "voxcpm2"


def _command(binary: Path, tool: str, request: dict, text_file: Path | None, engine: str = READ_ALOUD_ENGINE) -> list[str]:
    output = request["output"]
    # Files want accuracy over latency: ThunderTalk's most accurate model and
    # its best voice, not the dictation model or the Speak page's choice.
    # CLI version 1 (the first ThunderTalk with --cli) has no `best`.
    best = ["--engine", "best"] if int(request.get("thundertalk_cli", 1)) >= 2 else []
    if tool == "transcribe":
        options = request.get("options", {})
        command = [str(binary), "--cli", "transcribe", request["sources"][0], "--out", output,
                   "--format", options.get("format", "srt")]
        if options.get("speakers"):
            return [*command, "--speakers"]
        return [*command, *best]
    return [str(binary), "--cli", "speak", "--text-file", str(text_file), "--out", output,
            *(["--engine", engine] if engine else best)]


def run(request: dict, reporter: Reporter) -> None:
    binary = Path(request.get("thundertalk", ""))
    if not binary.is_file():
        raise reporter.error("speech_missing")
    tool = request["tool"]
    scratch = None
    text_file = None
    try:
        if tool == "speak":
            reporter.stage("stage_reading")
            source = Path(request["sources"][0])
            try:
                text = text_for_reading(source, request.get("source_suffix") or source.suffix)
            except ValueError:
                raise reporter.error("speech_unreadable") from None
            text = re.sub(r"[ \t]+", " ", text)
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
            if not text:
                raise reporter.error("speech_no_text")
            scratch = scratch_dir("speak-")
            text_file = scratch / "text.txt"
            text_file.write_text(text, encoding="utf-8")

        result: dict = {}

        def on_line(line: str) -> None:
            if not line.startswith("{"):
                return  # native libraries log to stderr; only JSON lines matter
            event = json.loads(line)
            kind = event.get("type")
            if kind == "stage" and event.get("label") in _STAGES:
                reporter.stage(_STAGES[event["label"]])
                reporter.indeterminate(event["label"] == "load_model")
            elif kind == "progress":
                reporter.indeterminate(False)
                reporter.progress(float(event.get("value", 0.0)))
            elif kind == "result":
                result.update(event)

        reporter.indeterminate(True)
        completed = run_tool(
            _command(binary, tool, request, text_file),
            timeout=TIMEOUT_SECONDS,
            timeout_message=reporter.text("tool_timeout", seconds=TIMEOUT_SECONDS),
            on_line=on_line,
        )
        if tool == "speak" and result.get("code") in {"tts_missing", "bad_args"}:
            # VoxCPM2 isn't downloaded in ThunderTalk here: use the best voice it has.
            result.clear()
            completed = run_tool(
                _command(binary, tool, request, text_file, engine=""),
                timeout=TIMEOUT_SECONDS,
                timeout_message=reporter.text("tool_timeout", seconds=TIMEOUT_SECONDS),
                on_line=on_line,
            )
        if result.get("ok"):
            reporter.progress(1.0)
            return
        code = str(result.get("code") or "")
        if code in _ERRORS:
            raise reporter.error(_ERRORS[code])
        detail = str(result.get("error") or "").strip()
        if not detail and completed.returncode != 0:
            detail = completed.stdout.strip().splitlines()[-1] if completed.stdout.strip() else ""
        message = reporter.text("speech_failed")
        if detail:
            message += ("：" if reporter.language == "zh" else ": ") + detail
        raise ConversionError(message)
    finally:
        if scratch is not None:
            remove_tree(scratch)
