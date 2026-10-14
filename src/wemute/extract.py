"""Any file → text a language model can read.

``extract(path)`` returns a :class:`Document` whose ``markdown`` keeps the
structure that carries meaning: headings, lists and tables in Word files,
tables and page numbers in PDFs (scanned pages are read with Apple Vision),
one section per sheet or slide, speakers and timestamps in recordings (through
ThunderTalk), headers and attachments in e-mails, and every member of an
archive. Used by ``Wemute --cli extract`` and tested file type by file type.
"""

from __future__ import annotations

import csv
import email
import email.policy
import html
import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from html.parser import HTMLParser
from pathlib import Path

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".log", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".xml",
    ".py", ".js", ".ts", ".sql", ".sh", ".r", ".tex", ".srt", ".vtt",
}
HTML_SUFFIXES = {".html", ".htm", ".xhtml"}
TEXTUTIL_SUFFIXES = {".doc", ".rtf", ".rtfd", ".odt", ".wordml", ".webarchive"}
OFFICE_SUFFIXES = {".pages", ".wps", ".numbers", ".key", ".ppt", ".odp", ".ods", ".xls"}
IMAGE_SUFFIXES = {
    ".jpg", ".jpeg", ".png", ".heic", ".heif", ".webp", ".gif", ".tif", ".tiff", ".bmp", ".avif",
}
AV_SUFFIXES = {
    ".mp3", ".m4a", ".wav", ".aac", ".flac", ".aiff", ".aif", ".ogg", ".opus", ".wma",
    ".mp4", ".mov", ".m4v", ".mkv", ".avi", ".webm",
}
ARCHIVE_SUFFIXES = {".zip", ".7z", ".tar", ".gz", ".tgz", ".bz2", ".xz", ".rar"}
EBOOK_SUFFIXES = {".mobi", ".azw", ".azw3", ".fb2"}

# A PDF page with less real text than this is treated as a scan and read with OCR.
SCANNED_PAGE_CHARS = 25
OCR_SCALE = 300 / 72


class ExtractError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass
class Document:
    source: str
    kind: str
    markdown: str
    title: str = ""
    meta: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    # Members of an archive or attachments of an e-mail.
    children: list["Document"] = field(default_factory=list)

    def plain(self) -> str:
        return markdown_to_plain(self.markdown)

    def to_dict(self) -> dict:
        data = asdict(self)
        data["text"] = self.plain()
        return data


@dataclass
class Options:
    ocr: str = "auto"  # auto | always | never
    speakers: bool = False
    language: str = "auto"
    max_rows: int = 5000  # per sheet; huge exports are cut with a note
    thundertalk: Path | None = None
    tools: dict[str, str] = field(default_factory=dict)  # installed components
    depth: int = 0


# ── entry point ──────────────────────────────────────────────────────────


def extract(path: Path, options: Options | None = None) -> Document:
    options = options or Options()
    path = Path(path)
    if not path.exists():
        raise ExtractError("not_found", f"File not found: {path}")
    if path.is_dir():
        return _folder(path, options)
    suffix = path.suffix.lower()
    if path.name.lower().endswith((".tar.gz", ".tar.bz2", ".tar.xz")):
        suffix = ".tar"
    handler = _handler(suffix, path)
    if handler is None:
        raise ExtractError("unsupported", f"Wemute can't read text from {suffix or 'this'} files yet")
    document = handler(path, options)
    document.source = str(path)
    if not document.title:
        document.title = path.stem
    return document


def _handler(suffix: str, path: Path):  # noqa: ANN202
    if suffix in TEXT_SUFFIXES:
        return _text
    if suffix in {".csv", ".tsv"}:
        return _delimited
    if suffix == ".json":
        return _json
    if suffix in HTML_SUFFIXES:
        return _html_file
    if suffix == ".docx":
        return _docx
    if suffix in {".xlsx", ".xlsm"}:
        return _xlsx
    if suffix == ".pptx":
        return _pptx
    if suffix == ".pdf":
        return _pdf
    if suffix == ".epub":
        return _epub
    if suffix == ".eml":
        return _eml
    if suffix in TEXTUTIL_SUFFIXES:
        return _textutil
    if suffix == ".svg":
        return _svg
    if suffix in OFFICE_SUFFIXES:
        return _via_office
    if suffix in EBOOK_SUFFIXES:
        return _via_calibre
    if suffix in IMAGE_SUFFIXES:
        return _image
    if suffix in AV_SUFFIXES:
        return _recording
    if suffix in ARCHIVE_SUFFIXES:
        return _archive
    if not suffix and _looks_textual(path):
        return _text
    return None


# ── plain text and data ─────────────────────────────────────────────────


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    for encoding in ("gb18030", "big5", "shift_jis", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _looks_textual(path: Path) -> bool:
    try:
        sample = path.read_bytes()[:4096]
    except OSError:
        return False
    return bool(sample) and b"\0" not in sample


def _text(path: Path, options: Options) -> Document:
    text = read_text(path)
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"} or not suffix:
        body = text
    elif suffix in {".txt", ".log"}:
        body = text
    else:
        language = {".py": "python", ".js": "javascript", ".ts": "typescript", ".sql": "sql",
                    ".sh": "sh", ".xml": "xml", ".yaml": "yaml", ".yml": "yaml", ".toml": "toml"}.get(suffix, "")
        body = f"```{language}\n{text.rstrip()}\n```"
    return Document(str(path), "text", body.strip() + "\n")


def md_table(rows: list[list[object]]) -> str:
    """A Markdown table; the first row is the header."""
    rows = [[_cell(value) for value in row] for row in rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
    return "\n".join(lines)


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if hasattr(value, "isoformat"):
        value = value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
        value = str(value).removesuffix(" 00:00:00")
    return str(value).replace("|", "\\|").replace("\r", "").replace("\n", "<br>").strip()


def _trim(rows: list[list[object]]) -> list[list[object]]:
    """Drop empty rows and trailing empty columns."""
    rows = [list(row) for row in rows if any(_cell(value) for value in row)]
    if not rows:
        return []
    width = max(
        (max((index + 1 for index, value in enumerate(row) if _cell(value)), default=0) for row in rows),
        default=0,
    )
    return [row[:width] for row in rows]


def _limited(rows: list[list[object]], options: Options, warnings: list[str], name: str) -> list[list[object]]:
    if len(rows) > options.max_rows + 1:
        warnings.append(f"{name}: first {options.max_rows} of {len(rows) - 1} rows")
        return rows[: options.max_rows + 1]
    return rows


def _delimited(path: Path, options: Options) -> Document:
    text = read_text(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else (csv.Sniffer().sniff(text[:4096], ",;\t|").delimiter if text.strip() else ",")
    rows = _trim(list(csv.reader(io.StringIO(text), delimiter=delimiter)))
    warnings: list[str] = []
    rows = _limited(rows, options, warnings, path.name)
    return Document(str(path), "table", md_table(rows) + "\n", meta={"rows": max(0, len(rows) - 1)}, warnings=warnings)


def _json(path: Path, options: Options) -> Document:
    data = json.loads(read_text(path))
    # A list of flat records reads best as a table.
    if isinstance(data, list) and data and all(isinstance(item, dict) for item in data):
        keys: list[str] = []
        for item in data:
            keys += [key for key in item if key not in keys]
        if all(not isinstance(value, (dict, list)) for item in data for value in item.values()):
            rows = [keys] + [[item.get(key) for key in keys] for item in data]
            warnings: list[str] = []
            rows = _limited(rows, options, warnings, path.name)
            return Document(str(path), "table", md_table(rows) + "\n", warnings=warnings)
    pretty = json.dumps(data, ensure_ascii=False, indent=2)
    return Document(str(path), "data", f"```json\n{pretty}\n```\n")


# ── HTML ─────────────────────────────────────────────────────────────────


class _HtmlToMarkdown(HTMLParser):
    """Headings, paragraphs, lists, links, emphasis and tables; scripts dropped."""

    BLOCKS = {"p", "div", "section", "article", "header", "footer", "main", "aside", "blockquote", "figure", "figcaption", "dd", "dt"}
    SKIP = {"script", "style", "noscript", "template", "svg", "head", "nav"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.out: list[str] = []
        self.skip = 0
        self.lists: list[str] = []
        self.counters: list[int] = []
        self.table: list[list[str]] | None = None
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.href: str | None = None
        self.link_text: list[str] = []
        self.title = ""
        self._in_title = False
        self.pre = 0

    def _write(self, text: str) -> None:
        if self.cell is not None:
            self.cell.append(text)
        elif self.href is not None:
            self.link_text.append(text)
        else:
            self.out.append(text)

    def _block(self) -> None:
        if self.cell is None:
            self.out.append("\n\n")

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        attrs = dict(attrs)
        if tag == "title":
            self._in_title = True
        if tag in self.SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            self._block()
            self._write("#" * int(tag[1]) + " ")
        elif tag in self.BLOCKS:
            self._block()
        elif tag == "br":
            self._write("\n")
        elif tag in {"ul", "ol"}:
            self.lists.append(tag)
            self.counters.append(0)
            self._block()
        elif tag == "li":
            depth = max(0, len(self.lists) - 1)
            if self.lists and self.lists[-1] == "ol":
                self.counters[-1] += 1
                marker = f"{self.counters[-1]}. "
            else:
                marker = "- "
            self._write("\n" + "  " * depth + marker)
        elif tag in {"strong", "b"}:
            self._write("**")
        elif tag in {"em", "i"}:
            self._write("*")
        elif tag == "pre":
            self.pre += 1
            self._block()
            self._write("```\n")
        elif tag == "a" and attrs.get("href") and self.cell is None:
            self.href = attrs["href"]
            self.link_text = []
        elif tag == "img" and attrs.get("alt"):
            self._write(f"[image: {attrs['alt']}]")
        elif tag == "table":
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):  # noqa: ANN001
        if tag == "title":
            self._in_title = False
        if tag in self.SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"} or tag in self.BLOCKS:
            self._block()
        elif tag in {"ul", "ol"} and self.lists:
            self.lists.pop()
            self.counters.pop()
            self._block()
        elif tag in {"strong", "b"}:
            self._write("**")
        elif tag in {"em", "i"}:
            self._write("*")
        elif tag == "pre":
            self.pre = max(0, self.pre - 1)
            self._write("\n```")
            self._block()
        elif tag == "a" and self.href is not None:
            text = "".join(self.link_text).strip()
            href, self.href = self.href, None
            if text and href.startswith(("http://", "https://")) and text != href:
                self._write(f"[{text}]({href})")
            else:
                self._write(text)
        elif tag in {"td", "th"} and self.cell is not None and self.row is not None:
            self.row.append(re.sub(r"\s+", " ", "".join(self.cell)).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None and self.table is not None:
            if self.row:
                self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            table, self.table = self.table, None
            if table:
                self.out.append("\n\n" + md_table(table) + "\n\n")

    def handle_data(self, data):  # noqa: ANN001
        if self._in_title:
            self.title += data
        if self.skip:
            return
        if self.pre:
            self._write(data)
        else:
            self._write(re.sub(r"\s+", " ", data))

    def markdown(self) -> str:
        text = "".join(self.out)
        text = re.sub(r"[ \t]+\n", "\n", text)
        text = re.sub(r"\n[ \t]+(?=[^\s-])", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip() + "\n"


def html_to_markdown(markup: str) -> tuple[str, str]:
    parser = _HtmlToMarkdown()
    parser.feed(markup)
    parser.close()
    return parser.markdown(), html.unescape(parser.title).strip()


def _html_file(path: Path, options: Options) -> Document:
    markdown, title = html_to_markdown(read_text(path))
    return Document(str(path), "web page", markdown, title=title)


# ── Office files ─────────────────────────────────────────────────────────


def _docx(path: Path, options: Options) -> Document:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    parts: list[str] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            level = re.match(r"(?:Heading|标题)\s*(\d)", style)
            if style == "Title":
                parts.append(f"# {text}")
            elif level:
                parts.append("#" * min(6, int(level.group(1)) + 1) + " " + text)
            elif "List" in style or child.find(".//{*}numPr") is not None:
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            rows = _trim(rows)
            if rows:
                parts.append(md_table(rows))
    core = document.core_properties
    meta = {key: value for key, value in {"author": core.author, "created": str(core.created or "")}.items() if value}
    markdown = _join_blocks(parts)
    return Document(str(path), "document", markdown, title=core.title or "", meta=meta)


def _join_blocks(parts: list[str]) -> str:
    """Blank lines between blocks, none between consecutive list items."""
    out: list[str] = []
    for part in parts:
        if out and not (part.startswith("- ") and out[-1].startswith("- ")):
            out.append("")
        out.append(part)
    return "\n".join(out).strip() + "\n"


def _xlsx(path: Path, options: Options) -> Document:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    sections: list[str] = []
    warnings: list[str] = []
    sheets = 0
    try:
        for sheet in workbook.worksheets:
            rows = _trim([list(row) for row in sheet.iter_rows(values_only=True)])
            if not rows:
                continue
            sheets += 1
            rows = _limited(rows, options, warnings, sheet.title)
            sections.append(f"## {sheet.title}\n\n{md_table(rows)}")
    finally:
        workbook.close()
    return Document(str(path), "spreadsheet", "\n\n".join(sections) + "\n", meta={"sheets": sheets}, warnings=warnings)


def _pptx(path: Path, options: Options) -> Document:
    from pptx import Presentation

    presentation = Presentation(str(path))
    sections: list[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        lines: list[str] = []
        for shape in sorted(slide.shapes, key=lambda item: (item.top or 0, item.left or 0)):
            if slide.shapes.title is not None and shape.shape_id == slide.shapes.title.shape_id:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append("  " * paragraph.level + f"- {text}")
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = _trim([[cell.text for cell in row.cells] for row in shape.table.rows])
                if rows:
                    lines.append("\n" + md_table(rows) + "\n")
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        section = f"## Slide {number}" + (f": {title}" if title else "")
        body = "\n".join(lines).strip()
        if body:
            section += "\n\n" + body
        if notes:
            section += f"\n\n> Notes: {notes}"
        sections.append(section)
    return Document(str(path), "slides", "\n\n".join(sections) + "\n", meta={"slides": len(sections)})


def _textutil(path: Path, options: Options) -> Document:
    """Word 97–2003, RTF, OpenDocument text and web archives, through macOS's textutil."""
    textutil = Path("/usr/bin/textutil")
    if not textutil.exists():
        raise ExtractError("unsupported", "textutil is not available on this system")
    result = subprocess.run(
        [str(textutil), "-convert", "html", "-stdout", str(path)],
        capture_output=True, timeout=120, check=False,
    )
    if result.returncode != 0 or not result.stdout:
        raise ExtractError("unreadable", f"Could not read {path.name}")
    markdown, title = html_to_markdown(result.stdout.decode("utf-8", errors="replace"))
    return Document(str(path), "document", markdown, title=title if title and title != path.name else "")


def _via_office(path: Path, options: Options) -> Document:
    """Pages, Numbers, Keynote, old Excel/PowerPoint: LibreOffice turns them into
    the modern format, which is then read like any other."""
    soffice = options.tools.get("libreoffice")
    if not soffice:
        raise ExtractError("component_missing", "Needs the Office component (Settings › Components › LibreOffice)")
    suffix = path.suffix.lower()
    target = {".numbers": "xlsx", ".ods": "xlsx", ".xls": "xlsx", ".key": "pptx", ".ppt": "pptx", ".odp": "pptx"}.get(suffix, "docx")
    with tempfile.TemporaryDirectory(prefix="wemute-extract-") as folder:
        profile = Path(folder) / "profile"
        result = subprocess.run(
            [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", target,
             "--outdir", folder, str(path)],
            capture_output=True, timeout=300, check=False,
        )
        converted = Path(folder) / f"{path.stem}.{target}"
        if result.returncode != 0 or not converted.exists():
            raise ExtractError("unreadable", f"LibreOffice could not open {path.name}")
        document = {"xlsx": _xlsx, "pptx": _pptx, "docx": _docx}[target](converted, options)
    return document


def _svg(path: Path, options: Options) -> Document:
    """Charts and diagrams: their title, description and every text label."""
    import xml.etree.ElementTree as ElementTree

    try:
        root = ElementTree.fromstring(path.read_bytes())
    except ElementTree.ParseError as exc:
        raise ExtractError("unreadable", f"Could not read the SVG: {exc}") from exc

    def local(tag: str) -> str:
        return tag.rsplit("}", 1)[-1]

    title = next((" ".join(e.itertext()).strip() for e in root.iter() if local(e.tag) == "title"), "")
    desc = next((" ".join(e.itertext()).strip() for e in root.iter() if local(e.tag) == "desc"), "")
    labels = [
        re.sub(r"\s+", " ", "".join(e.itertext())).strip()
        for e in root.iter()
        if local(e.tag) == "text"
    ]
    parts = [desc] if desc else []
    labels = [label for label in labels if label]
    if labels:
        parts.append("Text in the graphic:\n\n" + "\n".join(f"- {label}" for label in labels))
    meta = {key: root.get(key) for key in ("width", "height", "viewBox") if root.get(key)}
    markdown = "\n\n".join(parts) or "*(no text in the graphic)*"
    return Document(str(path), "vector", markdown + "\n", title=title, meta=meta)


# ── PDF ──────────────────────────────────────────────────────────────────


def _pdf(path: Path, options: Options) -> Document:
    import pymupdf as fitz

    fitz.TOOLS.mupdf_display_warnings(False)
    try:
        pdf = fitz.open(str(path))
    except Exception as exc:  # noqa: BLE001 - any parse failure
        raise ExtractError("unreadable", f"Could not open the PDF: {exc}") from exc
    if pdf.needs_pass:
        raise ExtractError("encrypted", "This PDF is password-protected")
    sections: list[str] = []
    scanned = 0
    tables = 0
    warnings: list[str] = []
    try:
        for index, page in enumerate(pdf):
            text_chars = len(page.get_text("text").strip())
            use_ocr = options.ocr == "always" or (options.ocr == "auto" and text_chars < SCANNED_PAGE_CHARS)
            if use_ocr and _has_ink(page):
                body = _ocr_pdf_page(page)
                if body:
                    scanned += 1
            else:
                body, found = _pdf_page_markdown(page)
                tables += found
            sections.append(f"## Page {index + 1}\n\n{body.strip()}" if body.strip() else f"## Page {index + 1}\n\n*(no text)*")
        metadata = {key: value for key, value in (pdf.metadata or {}).items() if value and key in {"title", "author", "subject", "creationDate"}}
        meta = {"pages": len(pdf), "ocr_pages": scanned, "tables": tables, **metadata}
        title = (pdf.metadata or {}).get("title") or ""
    finally:
        pdf.close()
    if scanned:
        warnings.append(f"{scanned} scanned page(s) read with OCR")
    return Document(str(path), "pdf", "\n\n".join(sections) + "\n", title=title, meta=meta, warnings=warnings)


def _has_ink(page) -> bool:  # noqa: ANN001
    return bool(page.get_images()) or bool(page.get_drawings()) or bool(page.get_text("text").strip())


def _pdf_page_markdown(page) -> tuple[str, int]:  # noqa: ANN001
    """Text blocks in reading order with tables as Markdown tables in place."""
    tables = []
    try:
        found = page.find_tables()
        tables = [table for table in found.tables if table.row_count >= 2 and table.col_count >= 2]
    except Exception:  # noqa: BLE001 - table detection is best effort
        tables = []
    boxes = [table.bbox for table in tables]

    def inside(rect) -> bool:  # noqa: ANN001
        x0, y0, x1, y1 = rect
        cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
        return any(bx0 <= cx <= bx1 and by0 <= cy <= by1 for bx0, by0, bx1, by1 in boxes)

    items: list[tuple[float, float, str]] = []
    for block in page.get_text("blocks", sort=True):
        x0, y0, x1, y1, text, _number, kind = block[:7]
        if kind != 0 or not text.strip() or inside((x0, y0, x1, y1)):
            continue
        items.append((y0, x0, _unwrap(text)))
    for table in tables:
        rows = _trim(table.extract())
        if rows:
            items.append((table.bbox[1], table.bbox[0], md_table(rows)))
    items.sort(key=lambda item: (round(item[0] / 4), item[1]))
    return "\n\n".join(text for _y, _x, text in items), len(tables)


def _unwrap(text: str) -> str:
    """Join lines a PDF wrapped mid-sentence; keep real line breaks."""
    lines = [line.strip() for line in text.strip().splitlines() if line.strip()]
    out: list[str] = []
    for line in lines:
        if out and not re.search(r"[。！？.!?:：;；]$", out[-1]) and not re.match(r"^([-•·*]|\d+[.)、])\s", line):
            joiner = "" if re.search(r"[\u4e00-\u9fff]$", out[-1]) and re.match(r"^[\u4e00-\u9fff]", line) else " "
            out[-1] = out[-1] + joiner + line
        else:
            out.append(line)
    return "\n".join(out)


def _ocr_pdf_page(page) -> str:  # noqa: ANN001
    from PIL import Image

    from wemute.engines import ocr

    pixmap = page.get_pixmap(matrix=__import__("pymupdf").Matrix(OCR_SCALE, OCR_SCALE), alpha=False)
    image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    return _paragraphs_markdown(ocr.paragraphs(ocr.recognize(image), *image.size))


def _paragraphs_markdown(paragraphs) -> str:  # noqa: ANN001
    blocks = []
    for paragraph in paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(f"### {text}" if paragraph.heading and "\n" not in text else text)
    return "\n\n".join(blocks)


# ── images ───────────────────────────────────────────────────────────────


def _image(path: Path, options: Options) -> Document:
    from wemute.engines import ocr
    from wemute.engines.vision_io import flatten, load_photo

    photo = load_photo(path)
    image = flatten(photo.image)
    meta: dict[str, object] = {"width": image.width, "height": image.height}
    taken = _exif_date(photo.image)
    if taken:
        meta["taken"] = taken
    markdown = ""
    warnings: list[str] = []
    if options.ocr != "never":
        markdown = _paragraphs_markdown(ocr.paragraphs(ocr.recognize(image), *image.size))
    codes = _codes(image)
    if codes:
        markdown += ("\n\n" if markdown else "") + "\n".join(f"- {kind}: {payload}" for kind, payload in codes)
        meta["codes"] = len(codes)
    if not markdown.strip():
        warnings.append("No text found in the image")
        markdown = "*(no text found)*"
    return Document(str(path), "image", markdown.strip() + "\n", meta=meta, warnings=warnings)


def _exif_date(image) -> str:  # noqa: ANN001
    try:
        exif = image.getexif()
        value = exif.get_ifd(0x8769).get(0x9003) or exif.get(0x0132)
        return str(value) if value else ""
    except Exception:  # noqa: BLE001
        return ""


def _codes(image) -> list[tuple[str, str]]:  # noqa: ANN001
    try:
        from wemute.engines.photo import read_codes

        return [(code.symbology, code.payload) for code in read_codes(image)]
    except Exception:  # noqa: BLE001 - codes are a bonus
        return []


# ── recordings (ThunderTalk) ─────────────────────────────────────────────


def _recording(path: Path, options: Options) -> Document:
    from wemute.integrations import thundertalk_binary, thundertalk_cli_version

    binary = options.thundertalk or thundertalk_binary()
    version = thundertalk_cli_version(Path(binary).parents[2]) if Path(binary).is_file() else 0
    if version < 1:
        raise ExtractError("thundertalk_missing", "Recordings need ThunderTalk with its command line (Settings › Integrations)")
    with tempfile.TemporaryDirectory(prefix="wemute-extract-") as folder:
        out = Path(folder) / "transcript.json"
        command = [str(binary), "--cli", "transcribe", str(path), "--out", str(out), "--format", "json"]
        if options.speakers:
            if version < 2:
                raise ExtractError("thundertalk_outdated", "Speaker labels need a newer ThunderTalk")
            command += ["--speakers"]
        elif version >= 2:
            command += ["--engine", "best"]
        if options.language != "auto":
            command += ["--language", options.language]
        result = subprocess.run(command, capture_output=True, text=True, timeout=6 * 3600, check=False)
        last = _last_json(result.stdout)
        if not out.exists() or not last.get("ok", False):
            code = str(last.get("code") or "failed")
            raise ExtractError(code, str(last.get("error") or "ThunderTalk could not transcribe this file"))
        data = json.loads(out.read_text(encoding="utf-8"))
    segments = data.get("segments") or []
    labelled = any(segment.get("speaker") for segment in segments)
    blocks: list[str] = []
    if labelled:
        turn_speaker, turn_start, turn_text = None, 0.0, []
        for segment in segments:
            speaker = segment.get("speaker") or ""
            if speaker != turn_speaker and turn_text:
                blocks.append(f"**{turn_speaker}** ({_clock(turn_start)})  \n" + _joined(turn_text))
                turn_text = []
            if not turn_text:
                turn_speaker, turn_start = speaker, float(segment.get("start", 0))
            turn_text.append(segment.get("text", ""))
        if turn_text:
            blocks.append(f"**{turn_speaker}** ({_clock(turn_start)})  \n" + _joined(turn_text))
    else:
        # Paragraph per ~30 s with a timestamp, so quotes can be found again.
        para_start, para = None, []
        for segment in segments:
            start = float(segment.get("start", 0))
            if para_start is None:
                para_start = start
            para.append(segment.get("text", ""))
            if start - para_start >= 30:
                blocks.append(f"`{_clock(para_start)}` " + _joined(para))
                para_start, para = None, []
        if para:
            blocks.append(f"`{_clock(para_start or 0)}` " + _joined(para))
    meta = {"duration": data.get("duration"), "engine": data.get("model") or data.get("engine"), "speakers": labelled}
    if not blocks:
        raise ExtractError("no_speech", "No speech was found in this recording")
    return Document(str(path), "recording", "\n\n".join(blocks) + "\n", meta=meta)


def _joined(texts: list[str]) -> str:
    """Segments of one turn or paragraph: a space between Latin words, none in Chinese."""
    out = ""
    for text in (t.strip() for t in texts):
        if not text:
            continue
        if out and re.search(r"[A-Za-z0-9,.!?;:]$", out) and re.match(r"[A-Za-z0-9]", text):
            out += " "
        out += text
    return out


def _clock(seconds: float) -> str:
    seconds = int(seconds)
    hours, rest = divmod(seconds, 3600)
    minutes, secs = divmod(rest, 60)
    return f"{hours}:{minutes:02d}:{secs:02d}" if hours else f"{minutes}:{secs:02d}"


def _last_json(stdout: str) -> dict:
    for line in reversed(stdout.splitlines()):
        if line.startswith("{"):
            try:
                return json.loads(line)
            except ValueError:
                continue
    return {}


# ── e-books, e-mail, archives, folders ───────────────────────────────────


def _epub(path: Path, options: Options) -> Document:
    with zipfile.ZipFile(path) as book:
        container = book.read("META-INF/container.xml").decode("utf-8", errors="replace")
        opf_path = re.search(r'full-path="([^"]+)"', container).group(1)
        opf = book.read(opf_path).decode("utf-8", errors="replace")
        base = opf_path.rsplit("/", 1)[0] + "/" if "/" in opf_path else ""
        manifest = dict(re.findall(r'<item\b[^>]*\bid="([^"]+)"[^>]*\bhref="([^"]+)"', opf))
        manifest.update({key: value for value, key in re.findall(r'<item\b[^>]*\bhref="([^"]+)"[^>]*\bid="([^"]+)"', opf)})
        spine = re.findall(r'<itemref\b[^>]*\bidref="([^"]+)"', opf)
        title = html.unescape(re.search(r"<dc:title[^>]*>(.*?)</dc:title>", opf, re.S).group(1)).strip() if "<dc:title" in opf else ""
        author = html.unescape(re.search(r"<dc:creator[^>]*>(.*?)</dc:creator>", opf, re.S).group(1)).strip() if "<dc:creator" in opf else ""
        chapters: list[str] = []
        for idref in spine:
            href = manifest.get(idref)
            if not href:
                continue
            try:
                markup = book.read(base + href).decode("utf-8", errors="replace")
            except KeyError:
                continue
            markdown, _ = html_to_markdown(markup)
            if markdown.strip():
                chapters.append(markdown.strip())
    meta = {"author": author} if author else {}
    return Document(str(path), "ebook", "\n\n---\n\n".join(chapters) + "\n", title=title, meta=meta | {"chapters": len(chapters)})


def _via_calibre(path: Path, options: Options) -> Document:
    calibre = options.tools.get("calibre")
    if not calibre:
        raise ExtractError("component_missing", "Needs the E-book component (Settings › Components › Calibre)")
    with tempfile.TemporaryDirectory(prefix="wemute-extract-") as folder:
        epub = Path(folder) / f"{path.stem}.epub"
        result = subprocess.run([calibre, str(path), str(epub)], capture_output=True, timeout=600, check=False)
        if result.returncode != 0 or not epub.exists():
            raise ExtractError("unreadable", f"Calibre could not open {path.name} (DRM-protected books can't be read)")
        return _epub(epub, options)


def _eml(path: Path, options: Options) -> Document:
    message = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    headers = [f"**{name}:** {message[name]}" for name in ("From", "To", "Cc", "Date", "Subject") if message[name]]
    body = message.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        content = body.get_content()
        text = html_to_markdown(content)[0] if body.get_content_type() == "text/html" else content.strip()
    children: list[Document] = []
    names: list[str] = []
    with tempfile.TemporaryDirectory(prefix="wemute-mail-") as folder:
        for part in message.iter_attachments():
            name = Path(part.get_filename() or "attachment").name
            names.append(name)
            target = Path(folder) / name
            payload = part.get_payload(decode=True) or b""
            target.write_bytes(payload)
            if options.depth < 3:
                children.append(_child(target, options, f"{path.name} › {name}"))
    markdown = "  \n".join(headers) + "\n\n" + text.strip() + "\n"
    if names:
        markdown += "\n**Attachments:** " + ", ".join(names) + "\n"
    return Document(str(path), "email", markdown, title=str(message["Subject"] or ""), children=children)


def _child(path: Path, options: Options, label: str) -> Document:
    nested = Options(**{**asdict(options), "depth": options.depth + 1})
    nested.thundertalk = options.thundertalk
    try:
        document = extract(path, nested)
    except ExtractError as exc:
        document = Document(label, "skipped", "", warnings=[str(exc)], meta={"error": exc.code})
    except Exception as exc:  # noqa: BLE001 - one bad member must not stop the rest
        document = Document(label, "skipped", "", warnings=[f"{type(exc).__name__}: {exc}"], meta={"error": "failed"})
    document.source = label
    return document


def _archive(path: Path, options: Options) -> Document:
    from wemute.engines.archive import extract as unpack
    from wemute.engines.common import Reporter

    if options.depth >= 3:
        raise ExtractError("too_deep", "Archive nested too deeply")
    folder = Path(tempfile.mkdtemp(prefix="wemute-archive-"))
    try:
        target = folder / "contents"
        unpack(path, target, Reporter("en", io.StringIO()), source_suffix=path.suffix.lower())
        members = sorted(item for item in target.rglob("*") if item.is_file() and not _hidden(item, target))
        children = [_child(item, options, f"{path.name} › {item.relative_to(target)}") for item in members]
    finally:
        shutil.rmtree(folder, ignore_errors=True)
    listing = "\n".join(f"- {Path(child.source).name if '›' not in child.source else child.source.split(' › ', 1)[1]}" for child in children)
    return Document(str(path), "archive", f"{len(children)} files:\n\n{listing}\n", children=children)


def _folder(path: Path, options: Options) -> Document:
    members = sorted(item for item in path.rglob("*") if item.is_file() and not _hidden(item, path))
    children = [_child(item, options, str(item)) for item in members]
    return Document(str(path), "folder", f"{len(children)} files\n", title=path.name, children=children)


def _hidden(item: Path, root: Path) -> bool:
    parts = item.relative_to(root).parts
    return any(part.startswith(".") or part == "__MACOSX" for part in parts)


# ── output ───────────────────────────────────────────────────────────────


def markdown_to_plain(markdown: str) -> str:
    text = re.sub(r"^#{1,6}\s*", "", markdown, flags=re.M)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`\n]+)`", r"[\1]", text)
    text = re.sub(r"^\|?-+(\|-+)*\|?$", "", text, flags=re.M)
    text = re.sub(r"^\|\s?|\s?\|$", "", text, flags=re.M)
    text = text.replace(" | ", "\t").replace("<br>", " ")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def flatten_documents(document: Document) -> list[Document]:
    """The document and every archive member / attachment, depth first."""
    out = [document]
    for child in document.children:
        out += flatten_documents(child)
    return out


def render_markdown(document: Document, level: int = 1) -> str:
    """One self-describing Markdown file: a small header, then the content.
    Archive members and attachments follow as sub-sections."""
    title = document.title or Path(document.source).name
    body = document.markdown.strip()
    if body.startswith("# "):  # the content's own title replaces ours
        first, _, body = body.partition("\n")
        title, body = first[2:].strip(), body.strip()
    header = [f"{'#' * level} {title}", "", f"- source: `{document.source}`", f"- type: {document.kind}"]
    for key, value in document.meta.items():
        if value not in ("", None, False) and not (value == 0 and key not in {"pages", "ocr_pages"}):
            header.append(f"- {key}: {value}")
    header += [f"- note: {warning}" for warning in document.warnings]
    if level > 1 and body:
        # Keep the content's headings below this document's own heading.
        body = re.sub(r"^(#{1,5}) ", lambda m: "#" * min(6, len(m.group(1)) + level - 1) + " ", body, flags=re.M)
    parts = ["\n".join(header)] + ([body] if body else [])
    parts += [render_markdown(child, min(6, level + 1)) for child in document.children]
    return "\n\n".join(parts).strip() + "\n"
