"""Real files for converter tests, generated on demand."""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tarfile
import zipfile
from pathlib import Path

TRUETYPE_FONT = Path("/System/Library/Fonts/Supplemental/Andale Mono.ttf")
CFF_FONT = Path("/System/Library/Fonts/Supplemental/NotoSansCanadianAboriginal-Regular.otf")


def make_docx(path: Path, paragraphs: int = 3) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Quarterly report 季度报告", 0)
    for index in range(paragraphs):
        document.add_paragraph(f"Paragraph {index + 1}: revenue grew steadily. 第 {index + 1} 段。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Sales"
    table.cell(1, 0).text, table.cell(1, 1).text = "North", "120"
    document.save(path)
    return path


def make_xlsx(path: Path) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Region", "Units", "Price"])
    sheet.append(["North", 120, 9.5])
    sheet.append(["南方", 80, 12.25])
    other = workbook.create_sheet("Notes")
    other.append(["Key", "Value"])
    other.append(["owner", "Ada"])
    workbook.save(path)
    return path


def make_single_sheet_xlsx(path: Path) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Region", "Units"])
    sheet.append(["North", 120])
    sheet.append(["South", 80])
    workbook.save(path)
    return path


def make_pptx(path: Path, slides: int = 3) -> Path:
    from pptx import Presentation

    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Slide {index + 1}"
        slide.placeholders[1].text = f"Point {index + 1} · 要点"
    presentation.save(path)
    return path


def make_pdf(path: Path, pages: int = 3) -> Path:
    """A text PDF drawn with Qt (real fonts, selectable text).

    Drawn in a child process: the test session may already own a non-GUI Qt app.
    """
    script = f"""
from PySide6.QtGui import QFont, QPainter, QPdfWriter, QPageSize
from wemute.engines.qt import ensure_gui
ensure_gui()
writer = QPdfWriter({str(path)!r})
writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
writer.setResolution(72)
painter = QPainter(writer)
painter.setFont(QFont("Helvetica", 18))
for index in range({pages}):
    if index:
        writer.newPage()
    painter.drawText(72, 100, f"Page {{index + 1}} of the sample report")
    painter.drawText(72, 140, "Wemute converts this locally.")
    painter.fillRect(72, 200, 300, 200, (index * 60) % 255)
painter.end()
"""
    subprocess.run([sys.executable, "-c", script], check=True, timeout=60)
    return path


def make_image_pdf(path: Path) -> Path:
    """A PDF holding one large photo-like image (for compression)."""
    from PIL import Image

    image = Image.effect_noise((2400, 1800), 40).convert("RGB")
    image = Image.merge("RGB", [image.getchannel(0), image.getchannel(0).rotate(90, expand=False), image.getchannel(0)])
    image.save(path, "PDF", resolution=150, quality=98)
    return path


def make_epub(path: Path, chapters: int = 2) -> Path:
    manifest = "".join(
        f'<item id="c{i}" href="c{i}.xhtml" media-type="application/xhtml+xml"/>' for i in range(chapters)
    )
    spine = "".join(f'<itemref idref="c{i}"/>' for i in range(chapters))
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        archive.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
            "</rootfiles></container>",
        )
        archive.writestr(
            "OEBPS/content.opf",
            '<?xml version="1.0" encoding="utf-8"?>'
            '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">'
            '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">wemute-test</dc:identifier>'
            "<dc:title>Sample Book</dc:title><dc:language>en</dc:language>"
            '<meta property="dcterms:modified">2026-01-01T00:00:00Z</meta></metadata>'
            f'<manifest><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>{manifest}</manifest>'
            f"<spine>{spine}</spine></package>",
        )
        archive.writestr(
            "OEBPS/nav.xhtml",
            '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml" '
            'xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Nav</title></head><body>'
            '<nav epub:type="toc"><ol>'
            + "".join(f'<li><a href="c{i}.xhtml">Chapter {i + 1}</a></li>' for i in range(chapters))
            + "</ol></nav></body></html>",
        )
        for index in range(chapters):
            archive.writestr(
                f"OEBPS/c{index}.xhtml",
                '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
                f"<head><title>Chapter {index + 1}</title></head><body><h1>Chapter {index + 1}</h1>"
                + "<p>It was a bright cold day in April. 这是一段测试文字。</p>" * 20
                + "</body></html>",
            )
    return path


def make_zip(path: Path, gbk_names: bool = False) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("notes/readme.txt", "hello\n")
        archive.writestr("notes/data/numbers.csv", "a,b\n1,2\n")
        archive.writestr("__MACOSX/notes/._readme.txt", "junk")
        archive.writestr("other.txt", "second top-level entry\n")
        if gbk_names:
            archive.writestr(_RawNameInfo("报告.txt".encode("gbk")), "中文文件名")
    return path


class _RawNameInfo(zipfile.ZipInfo):
    """Writes the name bytes as given, without the UTF-8 flag (like Windows zh)."""

    def __init__(self, raw: bytes) -> None:
        super().__init__(raw.decode("cp437"))
        self._raw = raw

    def _encodeFilenameFlags(self):  # noqa: N802, ANN202 - zipfile internals
        return self._raw, self.flag_bits & ~0x800


def make_wrapped_zip(path: Path) -> Path:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("photos/a.txt", "a")
        archive.writestr("photos/b.txt", "b")
    return path


def make_7z(path: Path) -> Path:
    import py7zr

    with py7zr.SevenZipFile(path, "w") as archive:
        archive.writestr("seven content", "inner/seven.txt")
        archive.writestr("second", "second.txt")
    return path


def make_tar_gz(path: Path, scratch: Path) -> Path:
    folder = scratch / "tarsrc"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "one.txt").write_text("one")
    (folder / "two.txt").write_text("two")
    with tarfile.open(path, "w:gz") as archive:
        archive.add(folder / "one.txt", arcname="one.txt")
        archive.add(folder / "two.txt", arcname="two.txt")
    return path


def make_csv(path: Path) -> Path:
    path.write_text("name,count,code\nAda,3,007\n林,12,042\n", encoding="utf-8")
    return path


def make_json(path: Path) -> Path:
    payload = {"items": [{"name": "Ada", "age": 36, "meta": {"team": "core"}}, {"name": "Lin", "age": 29}]}
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


def make_markdown(path: Path) -> Path:
    path.write_text(
        "# Release notes\n\nWemute now converts **documents**.\n\n"
        "| Kind | Engine |\n| --- | --- |\n| PDF | pikepdf |\n| 字体 | fontTools |\n\n"
        "```\ncode block\n```\n\n- one\n- two\n",
        encoding="utf-8",
    )
    return path


def make_html(path: Path) -> Path:
    path.write_text(
        "<html><head><title>Page</title><style>p{color:red}</style></head><body>"
        "<h1>Heading</h1><p>First <b>bold</b> paragraph.</p><ul><li>alpha</li><li>beta</li></ul>"
        "</body></html>",
        encoding="utf-8",
    )
    return path


def make_svg(path: Path) -> Path:
    path.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80" viewBox="0 0 120 80">'
        '<rect width="120" height="80" rx="12" fill="#2C4BF0"/>'
        '<circle cx="60" cy="40" r="22" fill="#FFFFFF"/></svg>',
        encoding="utf-8",
    )
    return path


def copy_font(path: Path, cff: bool = False) -> Path:
    # System fonts carry restricted file flags; copy only the bytes.
    shutil.copyfile(CFF_FONT if cff else TRUETYPE_FONT, path)
    return path


def make_png(path: Path, color: str = "teal", size: tuple[int, int] = (160, 120)) -> Path:
    from PIL import Image

    Image.new("RGB", size, color).save(path)
    return path


def make_jpeg(path: Path, size: tuple[int, int] = (200, 300)) -> Path:
    from PIL import Image

    Image.new("RGB", size, "orange").save(path, "JPEG", quality=90)
    return path
