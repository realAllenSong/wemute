"""Real LibreOffice / Calibre / Pandoc conversions through the controller.

Opt-in: install the components first (Settings → 组件, or ComponentManager),
then run with ``WEMUTE_E2E_COMPONENTS=1`` and the same ``WEMUTE_DATA_DIR``.
"""

from __future__ import annotations

import os
import uuid
from pathlib import Path

import docx
import pypdfium2 as pdfium
import pytest
from PySide6.QtCore import QCoreApplication, QEventLoop, QTimer

import doc_fixtures as fx
from wemute.components import ComponentManager
from wemute.components_qt import ComponentController
from wemute.controller import JobController
from wemute.domain import JobRecord, JobStatus
from wemute.persistence import QueueStore
from wemute.probe import probe_file
from wemute.queue_model import QueueModel

pytestmark = pytest.mark.skipif(
    os.environ.get("WEMUTE_E2E_COMPONENTS") != "1", reason="set WEMUTE_E2E_COMPONENTS=1"
)


def _pdf_text(path: Path) -> tuple[int, str]:
    document = pdfium.PdfDocument(path)
    return len(document), document[0].get_textpage().get_text_range()


def test_real_components_convert_documents_books_and_markdown(tmp_path: Path) -> None:
    QCoreApplication.instance() or QCoreApplication([])
    manager = ComponentManager()
    missing = {"libreoffice", "calibre", "pandoc"} - set(manager.tools())
    assert not missing, f"install these components first: {sorted(missing)}"
    sources = {
        "PDF:pptx": fx.make_pptx(tmp_path / "deck.pptx"),
        "PDF:docx": fx.make_docx(tmp_path / "report.docx"),
        "CSV": fx.make_single_sheet_xlsx(tmp_path / "sales.xlsx"),
        "PDF:epub": fx.make_epub(tmp_path / "book.epub"),
        "DOCX": fx.make_markdown(tmp_path / "notes.md"),
        "XLSX": fx.make_csv(tmp_path / "people.csv"),
    }
    jobs = {
        key: JobRecord(str(uuid.uuid4()), probe_file(path, create_preview=False), key.split(":")[0])
        for key, path in sources.items()
    }
    model = QueueModel(list(jobs.values()))
    controller = JobController(
        model, QueueStore(tmp_path / "queue.json"), components=ComponentController(manager)
    )
    controller._destination = str(tmp_path / "out")

    loop = QEventLoop()
    controller.busyChanged.connect(lambda: not controller.busy and loop.quit())
    QTimer.singleShot(600_000, loop.quit)
    controller.startAll()
    if controller.busy:
        loop.exec()

    for key, job in jobs.items():
        assert job.status == JobStatus.SUCCEEDED, f"{key}: {job.error}"
    pages, text = _pdf_text(Path(jobs["PDF:pptx"].output_path))
    assert pages == 3 and "Slide 1" in text
    pages, text = _pdf_text(Path(jobs["PDF:docx"].output_path))
    assert "Quarterly report" in text
    pages, text = _pdf_text(Path(jobs["PDF:epub"].output_path))
    assert pages >= 2 and "Chapter 1" in text
    csv_text = Path(jobs["CSV"].output_path).read_text(encoding="utf-8-sig")
    assert csv_text.splitlines()[0] == "Region,Units"
    document = docx.Document(jobs["DOCX"].output_path)
    assert document.paragraphs[0].style.name == "Heading 1"
    assert Path(jobs["XLSX"].output_path).suffix == ".xlsx"
    controller.shutdown()
