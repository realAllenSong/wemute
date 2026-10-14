"""`wemute extract`: every file type becomes text a language model can read.

Each test builds a real file with known content (Word, Excel, PowerPoint,
PDF with a table, a scanned PDF, a screenshot, e-mail with an attachment,
EPUB, archives, RTF/DOC via textutil…) and checks that the facts and the
structure survive. Recordings need ThunderTalk and are skipped without it.
"""

from __future__ import annotations

import json
import subprocess
import sys
import zipfile
from email.message import EmailMessage
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont

from wemute.extract import ExtractError, Options, extract, flatten_documents, render_markdown

FONT = "/System/Library/Fonts/Hiragino Sans GB.ttc"


def _font(size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(FONT, size)
    except OSError:
        pytest.skip("no CJK system font")


def _text_image(path: Path, lines: list[str], size=(1400, 900)) -> Path:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    y = 80
    for line in lines:
        draw.text((80, y), line, fill="black", font=_font(44))
        y += 90
    image.save(path)
    return path


# ── Office ───────────────────────────────────────────────────────────────


def test_docx_keeps_headings_lists_and_tables(tmp_path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("2026 年三季度业绩点评", level=1)
    document.add_paragraph("营收同比增长 18%，超出市场预期。")
    document.add_paragraph("毛利率承压", style="List Bullet")
    document.add_paragraph("库存周转改善", style="List Bullet")
    table = document.add_table(rows=3, cols=3)
    for row, values in enumerate([["指标", "Q3", "Q2"], ["营收（亿元）", "52.1", "47.3"], ["毛利率", "31.2%", "33.0%"]]):
        for col, value in enumerate(values):
            table.cell(row, col).text = value
    path = tmp_path / "report.docx"
    document.save(path)

    result = extract(path)
    md = result.markdown
    assert "## 2026 年三季度业绩点评" in md
    assert "- 毛利率承压\n- 库存周转改善" in md
    assert "| 营收（亿元） | 52.1 | 47.3 |" in md
    assert md.index("营收同比增长") < md.index("| 指标")


def test_xlsx_one_table_per_sheet(tmp_path: Path) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "利润表"
    sheet.append(["项目", "2025", "2026"])
    sheet.append(["营业收入", 180.5, 212.25])
    sheet.append(["净利润", 20, None])
    other = workbook.create_sheet("持仓")
    other.append(["代码", "名称", "权重"])
    other.append(["AAPL", "Apple", 0.12])
    path = tmp_path / "model.xlsx"
    workbook.save(path)

    md = extract(path).markdown
    assert "## 利润表" in md and "## 持仓" in md
    assert "| 营业收入 | 180.5 | 212.25 |" in md
    assert "| 净利润 | 20 |  |" in md
    assert "| AAPL | Apple | 0.12 |" in md


def test_pptx_slides_titles_bullets_and_notes(tmp_path: Path) -> None:
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "投资要点"
    slide.placeholders[1].text_frame.text = "AI 服务器需求强劲"
    slide.notes_slide.notes_text_frame.text = "强调订单能见度到 2027 年"
    second = deck.slides.add_slide(deck.slide_layouts[1])
    second.shapes.title.text = "风险"
    second.placeholders[1].text_frame.text = "汇率波动"
    path = tmp_path / "deck.pptx"
    deck.save(path)

    result = extract(path)
    assert "## Slide 1: 投资要点" in result.markdown
    assert "- AI 服务器需求强劲" in result.markdown
    assert "> Notes: 强调订单能见度到 2027 年" in result.markdown
    assert "## Slide 2: 风险" in result.markdown
    assert result.meta["slides"] == 2


@pytest.mark.parametrize("fmt", ["rtf", "doc", "odt"])
def test_textutil_formats(tmp_path: Path, fmt: str) -> None:
    if sys.platform != "darwin":
        pytest.skip("textutil is macOS only")
    source = tmp_path / "memo.html"
    source.write_text("<h1>会议纪要</h1><p>决定上调目标价至 220 美元。</p><table><tr><td>股票</td><td>评级</td></tr><tr><td>AAPL</td><td>买入</td></tr></table>", encoding="utf-8")
    target = tmp_path / f"memo.{fmt}"
    subprocess.run(["/usr/bin/textutil", "-inputencoding", "UTF-8", "-convert", fmt, str(source),
                    "-output", str(target)], check=True)
    md = extract(target).markdown
    assert "会议纪要" in md and "220 美元" in md
    assert "AAPL" in md and "买入" in md


# ── PDF ──────────────────────────────────────────────────────────────────


def test_pdf_text_pages_and_tables(tmp_path: Path) -> None:
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Quarterly Results", fontsize=20)
    page.insert_text((72, 110), "Revenue grew 18% year over year.", fontsize=11)
    # A ruled table.
    rows = [["Metric", "Q3", "Q2"], ["Revenue", "52.1", "47.3"], ["Margin", "31.2%", "33.0%"]]
    x0, y0, w, h = 72, 150, 120, 24
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = fitz.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 6, rect.y0 + 16), value, fontsize=10)
    second = pdf.new_page()
    second.insert_text((72, 72), "Guidance: revenue of 50 to 52 billion next quarter.", fontsize=11)
    path = tmp_path / "results.pdf"
    pdf.save(path)

    result = extract(path)
    md = result.markdown
    assert "## Page 1" in md and "## Page 2" in md
    assert "Revenue grew 18% year over year." in md
    assert "| Revenue | 52.1 | 47.3 |" in md
    assert "Guidance: revenue of 50 to 52 billion" in md
    assert result.meta["pages"] == 2 and result.meta["tables"] >= 1 and result.meta["ocr_pages"] == 0


@pytest.mark.skipif(sys.platform != "darwin", reason="Apple Vision")
def test_scanned_pdf_is_read_with_ocr(tmp_path: Path) -> None:
    image = _text_image(tmp_path / "scan.png", ["董事会决议公告", "同意回购公司股份 5 亿元"])
    path = tmp_path / "scan.pdf"
    Image.open(image).save(path, "PDF", resolution=150)
    result = extract(path)
    assert result.meta["ocr_pages"] == 1
    assert "董事会决议公告" in result.markdown
    assert "回购" in result.markdown and "5亿元" in result.markdown.replace(" ", "")


# ── images ───────────────────────────────────────────────────────────────


@pytest.mark.skipif(sys.platform != "darwin", reason="Apple Vision")
def test_screenshot_text(tmp_path: Path) -> None:
    path = _text_image(tmp_path / "shot.png", ["纳斯达克指数收涨 1.2%", "NVIDIA closed at $142.50"])
    result = extract(path)
    flat = result.markdown.replace(" ", "")
    assert "纳斯达克指数收涨1.2%" in flat
    assert "NVIDIA" in result.markdown and "142.50" in result.markdown
    assert result.meta["width"] == 1400


@pytest.mark.skipif(sys.platform != "darwin", reason="Apple Vision")
def test_blank_image_says_no_text(tmp_path: Path) -> None:
    path = tmp_path / "blank.png"
    Image.new("RGB", (400, 300), (40, 120, 200)).save(path)
    result = extract(path)
    assert result.warnings and "no text" in result.markdown


# ── web, data, text ──────────────────────────────────────────────────────


def test_html_to_markdown(tmp_path: Path) -> None:
    path = tmp_path / "news.html"
    path.write_text(
        "<html><head><title>市场快讯</title><script>var x=1</script></head><body><nav>菜单</nav>"
        "<h2>美联储维持利率不变</h2><p>声明称<b>通胀</b>仍高于目标。</p>"
        "<ul><li>点阵图显示年内降息两次</li><li>美元走弱</li></ul>"
        "<p><a href='https://example.com/fed'>原文</a></p>"
        "<table><tr><th>资产</th><th>涨跌</th></tr><tr><td>黄金</td><td>+1.1%</td></tr></table></body></html>",
        encoding="utf-8",
    )
    result = extract(path)
    md = result.markdown
    assert result.title == "市场快讯"
    assert "## 美联储维持利率不变" in md and "**通胀**" in md
    assert "- 点阵图显示年内降息两次" in md and "[原文](https://example.com/fed)" in md
    assert "| 黄金 | +1.1% |" in md
    assert "var x" not in md and "菜单" not in md


def test_csv_json_and_text(tmp_path: Path) -> None:
    csv_path = tmp_path / "trades.csv"
    csv_path.write_text("date,ticker,qty\n2026-10-01,AAPL,100\n2026-10-02,MSFT,-50\n", encoding="utf-8")
    assert "| 2026-10-02 | MSFT | -50 |" in extract(csv_path).markdown

    gbk = tmp_path / "gbk.csv"
    gbk.write_bytes("代码,名称\n600519,贵州茅台\n".encode("gb18030"))
    assert "贵州茅台" in extract(gbk).markdown

    records = tmp_path / "records.json"
    records.write_text(json.dumps([{"ticker": "AAPL", "pe": 31.5}, {"ticker": "TSLA", "pe": 70}]), encoding="utf-8")
    assert "| TSLA | 70 |" in extract(records).markdown

    nested = tmp_path / "nested.json"
    nested.write_text(json.dumps({"fund": {"name": "Alpha", "holdings": [1, 2]}}), encoding="utf-8")
    assert '"name": "Alpha"' in extract(nested).markdown

    note = tmp_path / "note.txt"
    note.write_text("电话会要点：订单饱满。", encoding="utf-8")
    assert "订单饱满" in extract(note).markdown


# ── e-mail, e-book, archives, folders ────────────────────────────────────


def test_email_with_attachment(tmp_path: Path) -> None:
    import openpyxl

    sheet_path = tmp_path / "attach.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["代码", "目标价"])
    workbook.active.append(["NVDA", 180])
    workbook.save(sheet_path)

    message = EmailMessage()
    message["From"] = "analyst@broker.com"
    message["To"] = "pm@fund.com"
    message["Subject"] = "NVDA 首次覆盖"
    message["Date"] = "Tue, 13 Oct 2026 09:00:00 +0800"
    message.set_content("首次覆盖给予买入评级，详见附件。")
    message.add_attachment(sheet_path.read_bytes(), maintype="application",
                           subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="估值.xlsx")
    path = tmp_path / "mail.eml"
    path.write_bytes(bytes(message))

    result = extract(path)
    assert "**From:** analyst@broker.com" in result.markdown
    assert "首次覆盖给予买入评级" in result.markdown and "估值.xlsx" in result.markdown
    assert result.children and "| NVDA | 180 |" in result.children[0].markdown
    assert "| NVDA | 180 |" in render_markdown(result)


def test_epub_chapters_in_order(tmp_path: Path) -> None:
    path = tmp_path / "book.epub"
    with zipfile.ZipFile(path, "w") as book:
        book.writestr("mimetype", "application/epub+zip")
        book.writestr("META-INF/container.xml",
                      '<container><rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>')
        book.writestr("OEBPS/content.opf",
                      '<package><metadata><dc:title>投资的本质</dc:title><dc:creator>某作者</dc:creator></metadata>'
                      '<manifest><item id="c1" href="c1.xhtml"/><item id="c2" href="c2.xhtml"/></manifest>'
                      '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>')
        book.writestr("OEBPS/c1.xhtml", "<html><body><h1>第一章 安全边际</h1><p>价格是你付出的。</p></body></html>")
        book.writestr("OEBPS/c2.xhtml", "<html><body><h1>第二章 复利</h1><p>时间是好生意的朋友。</p></body></html>")
    result = extract(path)
    assert result.title == "投资的本质" and result.meta["author"] == "某作者"
    md = result.markdown
    assert md.index("第一章 安全边际") < md.index("第二章 复利")
    assert "时间是好生意的朋友" in md


def test_zip_with_mixed_members_and_bad_member(tmp_path: Path) -> None:
    import docx

    inner = tmp_path / "inner.docx"
    document = docx.Document()
    document.add_paragraph("尽调纪要：渠道库存健康。")
    document.save(inner)
    path = tmp_path / "pack.zip"
    with zipfile.ZipFile(path, "w") as pack:
        pack.write(inner, "资料/纪要.docx")
        pack.writestr("资料/数据.csv", "a,b\n1,2\n")
        pack.writestr("资料/unknown.xyz", b"\x00\x01binary")
        pack.writestr("__MACOSX/._junk", b"x")
    result = extract(path)
    kinds = {child.source.split(" › ")[-1]: child.kind for child in result.children}
    assert kinds["纪要.docx"] == "document" and kinds["数据.csv"] == "table"
    assert any("渠道库存健康" in child.markdown for child in result.children)
    assert any(child.kind == "skipped" for child in result.children)
    assert not any("__MACOSX" in child.source for child in result.children)
    assert len(flatten_documents(result)) == 4


def test_folder_and_unsupported(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("alpha", encoding="utf-8")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / "b.md").write_text("# beta", encoding="utf-8")
    result = extract(tmp_path)
    assert {Path(child.source).name for child in result.children} == {"a.txt", "b.md"}
    weird = tmp_path / "x.bin"
    weird.write_bytes(b"\x00\x01\x02")
    with pytest.raises(ExtractError):
        extract(weird)


# ── recordings (ThunderTalk) ─────────────────────────────────────────────


def _thundertalk() -> Path:
    from wemute.integrations import thundertalk_binary, thundertalk_cli_version

    binary = thundertalk_binary()
    if thundertalk_cli_version() < 1:
        pytest.skip("ThunderTalk with --cli not installed")
    return binary


def _speak(binary: Path, text: str, voice: str, out: Path) -> None:
    result = subprocess.run([str(binary), "--cli", "speak", "--text", text, "--engine", voice.split(":")[0],
                             "--voice", voice, "--out", str(out)], capture_output=True, text=True, timeout=600)
    if not out.exists():
        pytest.skip(f"could not synthesize test speech: {result.stdout[-300:]}")


@pytest.mark.slow
def test_recording_transcript_and_speakers(tmp_path: Path) -> None:
    binary = _thundertalk()
    a = tmp_path / "a.wav"
    b = tmp_path / "b.wav"
    _speak(binary, "各位投资者大家好，本季度营收同比增长百分之十八。", "indextts:news-anchor-zh", a)
    _speak(binary, "请问管理层，下个季度的毛利率指引是多少？", "indextts:calm-male-zh", b)
    call = tmp_path / "call.m4a"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", str(a), "-i", str(b), "-filter_complex",
                    "[0][1]concat=n=2:v=0:a=1", str(call)], check=True)

    plain = extract(call)
    assert plain.kind == "recording" and "营收" in plain.markdown and "`0:00`" in plain.markdown

    labelled = extract(call, Options(speakers=True))
    assert labelled.meta["speakers"] is True
    assert "**S01**" in labelled.markdown and "**S02**" in labelled.markdown
    assert "毛利" in labelled.markdown


def test_cli_extract_stdout_and_json(tmp_path: Path) -> None:
    note = tmp_path / "n.txt"
    note.write_text("要点：维持买入。", encoding="utf-8")
    csv_path = tmp_path / "t.csv"
    csv_path.write_text("a,b\n1,2\n", encoding="utf-8")
    run = [sys.executable, "-m", "wemute", "--cli", "extract"]
    md = subprocess.run([*run, str(note), str(csv_path)], capture_output=True, text=True, timeout=120)
    assert md.returncode == 0 and "维持买入" in md.stdout and "| 1 | 2 |" in md.stdout
    js = subprocess.run([*run, "--format", "json", str(note), str(tmp_path / "missing.pdf")],
                        capture_output=True, text=True, timeout=120)
    lines = [json.loads(line) for line in js.stdout.splitlines()]
    assert js.returncode == 1
    assert lines[0]["type"] == "document" and "维持买入" in lines[0]["text"]
    assert lines[1]["type"] == "error" and lines[1]["code"] == "not_found"
    out = tmp_path / "out"
    written = subprocess.run([*run, "--out", str(out), str(note)], capture_output=True, text=True, timeout=120)
    assert written.returncode == 0 and (out / "n.md").read_text(encoding="utf-8").startswith("# n")
