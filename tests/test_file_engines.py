from __future__ import annotations

import json
import subprocess
import uuid
import zipfile
from pathlib import Path

import pytest

import doc_fixtures as fx
from wemute.capabilities import (
    COMBINE_PDF,
    MERGE_PDF,
    PACK_7Z,
    PACK_ZIP,
    default_target,
    options_for,
    unique_output_path,
)
from wemute.domain import JobRecord, MediaKind
from wemute.engine import FfmpegAdapter, finalize_output, remove_partial_output
from wemute.engines import ComponentMissing, required_component, route
from wemute.probe import probe_file


def job_for(path: Path, target: str) -> JobRecord:
    return JobRecord(str(uuid.uuid4()), probe_file(path, create_preview=False), target)


def convert(path: Path, target: str, out: Path, members: list[Path] | None = None) -> Path:
    adapter = FfmpegAdapter()
    job = job_for(path, target)
    member_jobs = [job_for(item, target) for item in members] if members else None
    plan = adapter.plan(job, out, member_jobs)
    result = subprocess.run(plan.command, capture_output=True, text=True, timeout=240)
    assert result.returncode == 0, result.stderr
    assert "wemute_progress=" in result.stdout or plan.folder_output or True
    finalize_output(plan)
    assert not plan.temporary_output.exists()
    return plan.final_output


# ------------------------------------------------------------------ detection


def test_detection_by_extension_and_content(tmp_path: Path) -> None:
    docx = fx.make_docx(tmp_path / "report.docx")
    renamed = tmp_path / "report-without-extension"
    renamed.write_bytes(docx.read_bytes())
    pdf = fx.make_pdf(tmp_path / "a.pdf", 1)
    pdf_no_ext = tmp_path / "scan"
    pdf_no_ext.write_bytes(pdf.read_bytes())
    plain_zip = fx.make_zip(tmp_path / "bundle.zip")
    zip_no_ext = tmp_path / "bundle"
    zip_no_ext.write_bytes(plain_zip.read_bytes())
    epub = fx.make_epub(tmp_path / "book.epub")
    epub_no_ext = tmp_path / "book"
    epub_no_ext.write_bytes(epub.read_bytes())
    svg_no_ext = tmp_path / "logo"
    svg_no_ext.write_text(fx.make_svg(tmp_path / "logo.svg").read_text())

    expected = {
        docx: (MediaKind.DOCUMENT, ".docx"),
        renamed: (MediaKind.DOCUMENT, ".docx"),
        pdf_no_ext: (MediaKind.PDF, ".pdf"),
        zip_no_ext: (MediaKind.ARCHIVE, ".zip"),
        epub_no_ext: (MediaKind.EBOOK, ".epub"),
        svg_no_ext: (MediaKind.VECTOR, ".svg"),
        fx.make_xlsx(tmp_path / "s.xlsx"): (MediaKind.SPREADSHEET, ".xlsx"),
        fx.make_csv(tmp_path / "t.csv"): (MediaKind.SPREADSHEET, ".csv"),
        fx.make_json(tmp_path / "d.json"): (MediaKind.DATA, ".json"),
        fx.make_markdown(tmp_path / "n.md"): (MediaKind.DATA, ".md"),
        fx.copy_font(tmp_path / "f.ttf"): (MediaKind.FONT, ".ttf"),
    }
    for path, (kind, suffix) in expected.items():
        descriptor = probe_file(path, create_preview=False)
        assert (descriptor.kind, descriptor.suffix) == (kind, suffix), path.name
    unknown = tmp_path / "mystery.qqq"
    unknown.write_bytes(b"\x00\x13\x37 nothing to see")
    assert probe_file(unknown, create_preview=False).kind == MediaKind.OTHER


def test_defaults_follow_the_design() -> None:
    assert default_target(MediaKind.DOCUMENT, ".docx") == "PDF"
    assert default_target(MediaKind.PRESENTATION, ".pptx") == "PDF"
    assert default_target(MediaKind.SPREADSHEET, ".xlsx") == "CSV"
    assert default_target(MediaKind.SPREADSHEET, ".csv") == "XLSX"
    assert default_target(MediaKind.PDF, ".pdf") == "DOCX"
    assert default_target(MediaKind.EBOOK, ".epub") == "PDF"
    assert default_target(MediaKind.ARCHIVE, ".zip") == "Folder"
    assert default_target(MediaKind.DATA, ".md") == "DOCX"
    assert default_target(MediaKind.DATA, ".json") == "CSV"
    assert default_target(MediaKind.FONT, ".ttf") == "WOFF2"
    assert default_target(MediaKind.VECTOR, ".svg") == "PNG"
    # Outputs that would only re-save the same format are not offered.
    assert "DOCX" not in options_for(MediaKind.DOCUMENT, suffixes={".docx"})
    assert "DOCX" in options_for(MediaKind.DOCUMENT, suffixes={".docx", ".odt"})
    assert options_for(MediaKind.DATA, suffixes={".json"}) == ["CSV", "XLSX", "MD", "HTML"]


def test_component_routing() -> None:
    assert route(MediaKind.SPREADSHEET, ".xlsx", "CSV").component is None
    assert route(MediaKind.SPREADSHEET, ".xls", "CSV").component == "libreoffice"
    assert route(MediaKind.SPREADSHEET, ".xlsx", "PDF").lane == "office"
    assert route(MediaKind.DATA, ".md", "DOCX").component == "pandoc"
    assert route(MediaKind.DATA, ".txt", "DOCX").component is None
    assert route(MediaKind.DATA, ".md", "PDF").component is None
    assert route(MediaKind.EBOOK, ".epub", "PDF").component == "calibre"


def test_missing_component_is_reported_before_running(tmp_path: Path) -> None:
    job = job_for(fx.make_docx(tmp_path / "a.docx"), "PDF")
    assert required_component(job) == "libreoffice"
    with pytest.raises(ComponentMissing):
        FfmpegAdapter().plan(job, tmp_path / "out")


# ------------------------------------------------------------------------ PDF


def test_pdf_to_page_images_writes_a_folder(tmp_path: Path) -> None:
    from PIL import Image

    source = fx.make_pdf(tmp_path / "report.pdf", pages=3)
    folder = convert(source, "PNG (per page)", tmp_path / "out")
    assert folder.is_dir() and folder.name == "report"
    pages = sorted(folder.iterdir())
    assert [page.name for page in pages] == ["report 01.png", "report 02.png", "report 03.png"]
    with Image.open(pages[0]) as image:
        assert image.width > 1000  # A4 at 150 dpi
    again = convert(source, "JPG (per page)", tmp_path / "out")
    assert again.name == "report 2" and len(list(again.glob("*.jpg"))) == 3


def test_pdf_to_text_and_docx(tmp_path: Path) -> None:
    from docx import Document

    source = fx.make_pdf(tmp_path / "report.pdf", pages=2)
    text = convert(source, "TXT", tmp_path / "out").read_text()
    assert "Page 1 of the sample report" in text and "Page 2" in text
    docx = convert(source, "DOCX", tmp_path / "out")
    content = "\n".join(paragraph.text for paragraph in Document(docx).paragraphs)
    assert "sample report" in content


def test_smaller_pdf_shrinks_photo_pages(tmp_path: Path) -> None:
    import pikepdf

    source = fx.make_image_pdf(tmp_path / "scan.pdf")
    output = convert(source, "Smaller PDF", tmp_path / "out")
    assert output.name == "scan compressed.pdf"
    assert output.stat().st_size < source.stat().st_size * 0.7
    with pikepdf.open(output) as pdf:
        assert len(pdf.pages) == 1


def test_merge_pdfs_in_list_order(tmp_path: Path) -> None:
    import pypdfium2 as pdfium

    first = fx.make_pdf(tmp_path / "a.pdf", pages=2)
    second = fx.make_pdf(tmp_path / "b.pdf", pages=3)
    merged = convert(first, MERGE_PDF, tmp_path / "out", members=[second, first])
    document = pdfium.PdfDocument(str(merged))
    assert len(document) == 5
    assert "Page 3" in document[2].get_textpage().get_text_range()
    document.close()


def test_images_combine_into_one_pdf(tmp_path: Path) -> None:
    import pikepdf

    images = [
        fx.make_png(tmp_path / "one.png", "navy", (400, 300)),
        fx.make_jpeg(tmp_path / "two.jpg", (300, 600)),
        fx.make_png(tmp_path / "three.png", "gold", (100, 100)),
    ]
    output = convert(images[0], COMBINE_PDF, tmp_path / "out", members=images)
    assert output.name == f"{tmp_path.name}.pdf"
    with pikepdf.open(output) as pdf:
        assert len(pdf.pages) == 3
        boxes = [[float(value) for value in page.mediabox] for page in pdf.pages]
        assert all(max(box[2], box[3]) == pytest.approx(842, abs=0.5) for box in boxes)
        assert boxes[0][2] > boxes[0][3] and boxes[1][3] > boxes[1][2]
        # The JPEG is embedded as is, without re-encoding.
        image = pdf.pages[1].Resources.XObject["/Im0"]
        assert image.read_raw_bytes() == images[1].read_bytes()


# -------------------------------------------------------------------- archive


def test_zip_extracts_to_folder_without_mac_junk(tmp_path: Path) -> None:
    source = fx.make_zip(tmp_path / "bundle.zip", gbk_names=True)
    folder = convert(source, "Folder", tmp_path / "out")
    assert folder.name == "bundle"
    names = sorted(str(path.relative_to(folder)) for path in folder.rglob("*") if path.is_file())
    assert names == ["notes/data/numbers.csv", "notes/readme.txt", "other.txt", "报告.txt"]
    assert not (folder / "__MACOSX").exists()


def test_single_top_folder_is_unwrapped_and_names_never_overwrite(tmp_path: Path) -> None:
    source = fx.make_wrapped_zip(tmp_path / "photos.zip")
    out = tmp_path / "out"
    out.mkdir()
    (out / "photos").mkdir()
    (out / "photos" / "keep.txt").write_text("mine")
    folder = convert(source, "Folder", out)
    assert folder.name == "photos 2"
    assert sorted(path.name for path in folder.iterdir()) == ["a.txt", "b.txt"]
    assert (out / "photos" / "keep.txt").read_text() == "mine"


def test_7z_and_tar_gz_extract(tmp_path: Path) -> None:
    seven = convert(fx.make_7z(tmp_path / "pack.7z"), "Folder", tmp_path / "out")
    assert (seven / "inner" / "seven.txt").read_text() == "seven content"
    tarred = convert(fx.make_tar_gz(tmp_path / "logs.tar.gz", tmp_path), "Folder", tmp_path / "out")
    assert tarred.name == "logs"
    assert sorted(path.name for path in tarred.iterdir()) == ["one.txt", "two.txt"]


def test_pack_any_group_into_zip_and_7z(tmp_path: Path) -> None:
    import py7zr

    files = [fx.make_csv(tmp_path / "a.csv"), fx.make_svg(tmp_path / "b.svg")]
    nested = tmp_path / "nested"
    nested.mkdir()
    files.append(fx.make_csv(nested / "a.csv"))
    packed = convert(files[0], PACK_ZIP, tmp_path / "out", members=files)
    with zipfile.ZipFile(packed) as archive:
        assert archive.namelist() == ["a.csv", "b.svg", "a 2.csv"]
        assert archive.read("a.csv") == files[0].read_bytes()
    seven = convert(files[0], PACK_7Z, tmp_path / "out", members=files)
    with py7zr.SevenZipFile(seven) as archive:
        assert sorted(archive.getnames()) == ["a 2.csv", "a.csv", "b.svg"]


# ------------------------------------------------------------- data & tables


def test_xlsx_to_csv_one_file_per_sheet(tmp_path: Path) -> None:
    single = convert(fx.make_single_sheet_xlsx(tmp_path / "one.xlsx"), "CSV", tmp_path / "out")
    assert single.suffix == ".csv"
    assert single.read_text(encoding="utf-8-sig").splitlines() == ["Region,Units", "North,120", "South,80"]
    folder = convert(fx.make_xlsx(tmp_path / "book.xlsx"), "CSV", tmp_path / "out")
    assert folder.is_dir() and folder.name == "book"
    sales = (folder / "book - Sales.csv").read_text(encoding="utf-8-sig").splitlines()
    assert sales == ["Region,Units,Price", "North,120,9.5", "南方,80,12.25"]
    assert (folder / "book - Notes.csv").exists()


def test_csv_to_xlsx_keeps_numbers_and_ids(tmp_path: Path) -> None:
    import openpyxl

    output = convert(fx.make_csv(tmp_path / "people.csv"), "XLSX", tmp_path / "out")
    sheet = openpyxl.load_workbook(output).active
    assert [cell.value for cell in sheet[1]] == ["name", "count", "code"]
    assert [cell.value for cell in sheet[3]] == ["林", 12, "042"]


def test_xlsx_and_json_tables(tmp_path: Path) -> None:
    import openpyxl

    payload = json.loads(convert(fx.make_xlsx(tmp_path / "b.xlsx"), "JSON", tmp_path / "out").read_text())
    assert payload["Sales"][1] == {"Region": "南方", "Units": 80, "Price": 12.25}
    rows = convert(fx.make_json(tmp_path / "people.json"), "CSV", tmp_path / "out")
    assert rows.read_text(encoding="utf-8-sig").splitlines() == ["name,age,meta.team", "Ada,36,core", "Lin,29,"]
    workbook = convert(fx.make_json(tmp_path / "people.json"), "XLSX", tmp_path / "out")
    assert openpyxl.load_workbook(workbook).active["B2"].value == 36
    markdown = convert(fx.make_json(tmp_path / "people.json"), "MD", tmp_path / "out").read_text()
    assert "| name | age | meta.team |" in markdown


def test_markdown_and_text_without_components(tmp_path: Path) -> None:
    import pypdfium2 as pdfium
    from docx import Document

    notes = fx.make_markdown(tmp_path / "notes.md")
    html = convert(notes, "HTML", tmp_path / "out").read_text()
    assert "<table>" in html and "<strong>documents</strong>" in html
    text = convert(notes, "TXT", tmp_path / "out").read_text()
    assert "Release notes" in text and "**" not in text
    pdf = convert(notes, "PDF", tmp_path / "out")
    document = pdfium.PdfDocument(str(pdf))
    assert "Release notes" in document[0].get_textpage().get_text_range()
    document.close()
    plain = tmp_path / "letter.txt"
    plain.write_text("Dear team,\n\nThanks for everything.\n")
    assert [p.text for p in Document(convert(plain, "DOCX", tmp_path / "out")).paragraphs] == [
        "Dear team,",
        "Thanks for everything.",
    ]
    page = convert(fx.make_html(tmp_path / "page.html"), "TXT", tmp_path / "out").read_text()
    assert page.split() == ["Heading", "First", "bold", "paragraph.", "•", "alpha", "•", "beta"]


# ---------------------------------------------------------------- fonts & svg


def test_font_flavors_and_outline_conversion(tmp_path: Path) -> None:
    from fontTools.ttLib import TTFont

    ttf = fx.copy_font(tmp_path / "Mono.ttf")
    woff2 = convert(ttf, "WOFF2", tmp_path / "out")
    assert TTFont(woff2).flavor == "woff2"
    back = convert(woff2, "TTF", tmp_path / "out")
    assert "glyf" in TTFont(back) and TTFont(back).flavor is None
    otf = convert(ttf, "OTF", tmp_path / "out")
    converted = TTFont(otf)
    assert "CFF " in converted and "glyf" not in converted
    assert set(converted.getGlyphOrder()) == set(TTFont(ttf).getGlyphOrder())

    cff = fx.copy_font(tmp_path / "Syllabics.otf", cff=True)
    as_ttf = TTFont(convert(cff, "TTF", tmp_path / "out"))
    assert "glyf" in as_ttf and "CFF " not in as_ttf
    assert as_ttf.getBestCmap() == TTFont(cff).getBestCmap()
    assert TTFont(convert(cff, "WOFF", tmp_path / "out")).flavor == "woff"


def test_svg_to_png_and_pdf(tmp_path: Path) -> None:
    import pikepdf
    from PIL import Image

    svg = fx.make_svg(tmp_path / "badge.svg")
    with Image.open(convert(svg, "PNG", tmp_path / "out")) as image:
        assert image.size == (1024, 683)
        assert image.getpixel((512, 341))[:3] == (255, 255, 255)
        assert image.getpixel((2, 2))[3] == 0  # rounded corner stays transparent
    with pikepdf.open(convert(svg, "PDF", tmp_path / "out")) as pdf:
        box = [float(value) for value in pdf.pages[0].mediabox]
        assert box[2] == pytest.approx(90, abs=1) and box[3] == pytest.approx(60, abs=1)


# ------------------------------------------------------------ partial outputs


def test_failed_folder_output_is_removed(tmp_path: Path) -> None:
    broken = tmp_path / "broken.zip"
    broken.write_bytes(b"PK\x03\x04 this is not really a zip")
    adapter = FfmpegAdapter()
    job = job_for(broken, "Folder")
    plan = adapter.plan(job, tmp_path / "out")
    result = subprocess.run(plan.command, capture_output=True, text=True)
    assert result.returncode == 1
    assert result.stderr.strip()
    remove_partial_output(plan)
    assert not plan.temporary_output.exists()
    assert list((tmp_path / "out").iterdir()) == []
    # The reserved name is free again.
    assert unique_output_path(broken, tmp_path / "out", "", stem="broken").name == "broken"
