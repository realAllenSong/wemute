"""Apple Vision / Core Image tools, run through the real document worker.

``fixtures/vision/nasa-astronauts-2013.jpg`` is a NASA photo (Mark Sowa, NASA JSC,
2013; public domain), downscaled. The other inputs are drawn here so the tests
know exactly what Vision should find.
"""

from __future__ import annotations

import csv
import subprocess
import uuid
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageDraw, ImageFilter, ImageFont

import doc_fixtures as fx
from wemute.domain import JobRecord, MediaKind
from wemute.engine import FfmpegAdapter, finalize_output
from wemute.probe import create_thumbnail, probe_file
from wemute.tools import TOOLS, tool_for, tools_for, vision_available

pytestmark = pytest.mark.skipif(not vision_available(), reason="needs Apple Vision (macOS 14+)")

FIXTURES = Path(__file__).parent / "fixtures" / "vision"
NASA = FIXTURES / "nasa-astronauts-2013.jpg"
CJK_FONT = Path("/System/Library/Fonts/Hiragino Sans GB.ttc")


# ── helpers ──────────────────────────────────────────────────────────────


def _job(path: Path, target: str) -> JobRecord:
    return JobRecord(str(uuid.uuid4()), probe_file(path, create_preview=False), target)


def run_tool(
    path: Path, tool: str, out: Path, members: list[Path] | None = None, language: str = "en"
) -> tuple[Path | None, subprocess.CompletedProcess]:
    """Plan ``tool`` like the app does and run its worker; (output, process)."""
    job = _job(path, tool)
    member_jobs = [_job(item, tool) for item in members] if members else None
    adapter = FfmpegAdapter()
    adapter.language = language
    plan = adapter.plan(job, out, member_jobs)
    result = subprocess.run(plan.command, capture_output=True, text=True, timeout=300)
    if result.returncode != 0:
        return None, result
    finalize_output(plan)
    return plan.final_output, result


def ok(path: Path, tool: str, out: Path, members: list[Path] | None = None, language: str = "en") -> Path:
    output, result = run_tool(path, tool, out, members, language)
    assert output is not None, result.stderr
    assert "wemute_stage=" in result.stdout
    return output


def failure(path: Path, tool: str, out: Path) -> str:
    output, result = run_tool(path, tool, out)
    assert output is None, "expected a readable failure"
    return result.stderr.strip().splitlines()[-1]


def _font(size: int) -> ImageFont.FreeTypeFont:
    if not CJK_FONT.exists():
        pytest.skip("needs the Hiragino Sans GB system font")
    return ImageFont.truetype(str(CJK_FONT), size)


def text_image(path: Path, lines: list[str], size: tuple[int, int] = (1400, 700)) -> Path:
    image = Image.new("RGB", size, (250, 250, 252))
    draw = ImageDraw.Draw(image)
    font = _font(40)
    for index, line in enumerate(lines):
        draw.text((70, 70 + index * 80), line, font=font, fill=(25, 30, 45))
    image.save(path)
    return path


def page_photo(path: Path) -> Path:
    """A printed page photographed at an angle on a dark, textured table."""
    rng = np.random.default_rng(2)
    table = Image.fromarray(np.clip(rng.normal(70, 18, (1500, 2000, 3)), 0, 255).astype(np.uint8))
    table = table.filter(ImageFilter.GaussianBlur(2))
    page = Image.new("RGB", (850, 1100), (250, 248, 240))
    draw = ImageDraw.Draw(page)
    font = _font(40)
    for index in range(14):
        draw.text((70, 80 + index * 68), f"第 {index + 1} 行 会议纪要 Meeting notes", font=font, fill=(30, 30, 40))
    # A pinhole camera (f = 1700 px) looking at the page tilted 28° and 12°.
    corners = [(587, 166), (1534, 180), (1413, 1198), (705, 1101)]
    source = [(0, 0), (850, 0), (850, 1100), (0, 1100)]
    rows, values = [], []
    for (x, y), (u, v) in zip(corners, source):
        rows += [[x, y, 1, 0, 0, 0, -u * x, -u * y], [0, 0, 0, x, y, 1, -v * x, -v * y]]
        values += [u, v]
    coefficients = tuple(np.linalg.solve(np.array(rows, float), np.array(values, float)))
    warped = page.transform(table.size, Image.Transform.PERSPECTIVE, coefficients, Image.Resampling.BICUBIC)
    mask = Image.new("L", page.size, 255).transform(
        table.size, Image.Transform.PERSPECTIVE, coefficients, Image.Resampling.BICUBIC
    )
    table.paste(warped, (0, 0), mask)
    table.save(path, quality=90)
    return path


def tilted_sea(path: Path, degrees: float = 4.0) -> Path:
    width, height = 1600, 1066
    rows = np.linspace(0, 1, height)[:, None, None]
    sky = np.array([150, 190, 235])[None, None, :] * (1 - 0.35 * rows)
    sea = np.array([25, 70, 120])[None, None, :] + np.random.default_rng(1).normal(0, 8, (height, width, 3))
    pixels = np.where((np.arange(height) < height * 0.55)[:, None, None], sky, sea)
    image = Image.fromarray(np.clip(pixels, 0, 255).astype(np.uint8))
    image = image.rotate(-degrees, resample=Image.Resampling.BICUBIC).crop((80, 60, width - 80, height - 60))
    image.save(path, quality=92)
    return path


def qr_code(path: Path, message: str) -> Path:
    """Drawn by Core Image's own generator."""
    import Quartz
    from Foundation import NSData

    from wemute.engines.vision_io import render_ci

    generator = Quartz.CIFilter.filterWithName_("CIQRCodeGenerator")
    generator.setValue_forKey_(NSData.dataWithBytes_length_(message.encode(), len(message.encode())), "inputMessage")
    generator.setValue_forKey_("M", "inputCorrectionLevel")
    code = generator.outputImage().imageByApplyingTransform_(Quartz.CGAffineTransformMakeScale(12, 12))
    image = render_ci(code).convert("RGB")
    canvas = Image.new("RGB", (image.width + 120, image.height + 120), "white")
    canvas.paste(image, (60, 60))
    canvas.save(path)
    return path


def plain(path: Path, color: tuple[int, int, int] = (128, 140, 150)) -> Path:
    Image.new("RGB", (800, 600), color).save(path)
    return path


# ── menus ────────────────────────────────────────────────────────────────


def test_tools_are_offered_by_kind_and_count(monkeypatch, tmp_path: Path) -> None:  # noqa: ANN001
    # Speech tools depend on ThunderTalk; keep this test about Vision.
    monkeypatch.setenv("WEMUTE_THUNDERTALK_APP", str(tmp_path / "none.app"))
    single = {tool.key for tool in tools_for(MediaKind.IMAGE, 1)}
    group = {tool.key for tool in tools_for(MediaKind.IMAGE, 3)}
    assert "Blur faces" in single and "Blurry & duplicates" not in single
    assert {"Blurry & duplicates", "One searchable PDF", "Codes CSV"} <= group
    assert {tool.key for tool in tools_for(MediaKind.PDF, 1)} == {"Text (OCR)", "Searchable PDF"}
    assert not tools_for(MediaKind.AUDIO, 2)
    assert all(tool_for(tool.key) is tool for tool in TOOLS if tool.needs == "vision")


def test_image_menu_has_sections_columns_and_variant_chips(tmp_path: Path) -> None:
    from PySide6.QtCore import QCoreApplication

    from wemute.components import ComponentManager
    from wemute.components_qt import ComponentController
    from wemute.controller import JobController
    from wemute.persistence import QueueStore
    from wemute.queue_model import QueueModel

    QCoreApplication.instance() or QCoreApplication([])
    jobs = [_job(fx.make_jpeg(tmp_path / f"photo {index}.jpg"), "Blur faces") for index in range(2)]
    model = QueueModel(jobs)
    model.set_language("zh")
    components = ComponentController(ComponentManager(tmp_path / "components", []))
    controller = JobController(model, QueueStore(tmp_path / "queue.json"), components=components)
    try:
        (group,) = controller.groups
        assert (group["lead"], group["verb"], group["targetLabel"]) == ("给", "", "人脸打码")
        assert group["extra"] == ""
        entries = group["entries"]
        headers = [entry["header"] for entry in entries if entry.get("header")]
        assert headers[:3] == ["格式", "照片工具", "识别与整理"]
        assert sum(entry.get("columnBreak") is True for entry in entries) == 1
        grid = next(entry for entry in entries if entry.get("grid"))
        assert [chip["label"] for chip in grid["grid"]][:3] == ["WebP", "JPEG", "PNG"]
        faces = next(entry for entry in entries if entry.get("label") == "人脸打码")
        assert faces["checked"] and [chip["label"] for chip in faces["variants"]] == ["模糊", "马赛克"]
        assert faces["variants"][0]["checked"]
        assert "Blurry & duplicates" in group["options"]
    finally:
        controller.shutdown()


# ── OCR ──────────────────────────────────────────────────────────────────


LINES = ["季度报告 Quarterly Report 2026", "本季度营收增长 18%，主要来自海外市场。", "Next steps: expand the team."]


def test_ocr_image_to_text_and_word(tmp_path: Path) -> None:
    import docx

    source = text_image(tmp_path / "report.png", LINES)
    text = ok(source, "Text (OCR)", tmp_path / "out").read_text(encoding="utf-8")
    assert "季度报告" in text and "Quarterly Report" in text
    assert "营收增长" in text and "expand the team" in text
    document = docx.Document(ok(source, "Word (OCR)", tmp_path / "out"))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    assert "季度报告" in paragraphs[0] and "expand the team" in paragraphs[-1]


def test_ocr_searchable_pdf_has_an_invisible_text_layer(tmp_path: Path) -> None:
    import pypdfium2 as pdfium

    source = text_image(tmp_path / "report.png", LINES)
    output = ok(source, "Searchable PDF", tmp_path / "out")
    document = pdfium.PdfDocument(str(output))
    try:
        page = document[0]
        text = page.get_textpage().get_text_range()
        assert "季度报告" in text and "Quarterly Report" in text
        # The page still looks like the photo: the text layer draws nothing.
        rendered = np.asarray(page.render(scale=1).to_pil().convert("L"))
        original = np.asarray(Image.open(source).convert("L").resize(rendered.shape[::-1]))
        assert np.abs(rendered.astype(int) - original.astype(int)).mean() < 6
    finally:
        document.close()


def test_ocr_group_makes_one_searchable_pdf(tmp_path: Path) -> None:
    import pypdfium2 as pdfium

    first = text_image(tmp_path / "a.png", ["第一页 First page"])
    second = text_image(tmp_path / "b.png", ["第二页 Second page"])
    output = ok(first, "One searchable PDF", tmp_path / "out", members=[first, second])
    document = pdfium.PdfDocument(str(output))
    try:
        texts = [document[index].get_textpage().get_text_range() for index in range(len(document))]
    finally:
        document.close()
    assert len(texts) == 2 and "First page" in texts[0] and "Second page" in texts[1]


def test_pdf_ocr_keeps_digital_pages_and_reads_scanned_ones(tmp_path: Path) -> None:
    import pikepdf
    import pypdfium2 as pdfium

    scanned = Image.open(text_image(tmp_path / "scan.png", ["扫描的页面 Scanned page"])).convert("RGB")
    scanned_pdf = tmp_path / "scan.pdf"
    scanned.save(scanned_pdf, resolution=150)
    digital = fx.make_pdf(tmp_path / "digital.pdf", pages=1)
    combined = tmp_path / "mixed.pdf"
    with pikepdf.new() as pdf, pikepdf.open(scanned_pdf) as a, pikepdf.open(digital) as b:
        pdf.pages.extend(a.pages)
        pdf.pages.extend(b.pages)
        pdf.save(combined)
    with pdfium.PdfDocument(str(digital)) as original:
        digital_text = original[0].get_textpage().get_text_range()

    output = ok(combined, "Searchable PDF", tmp_path / "out")
    assert output.name == "mixed searchable.pdf"
    with pdfium.PdfDocument(str(output)) as document:
        assert "Scanned page" in document[0].get_textpage().get_text_range()
        assert document[1].get_textpage().get_text_range() == digital_text
    text = ok(combined, "Text (OCR)", tmp_path / "out").read_text(encoding="utf-8")
    assert "扫描的页面" in text


def test_ocr_reports_a_blank_image(tmp_path: Path) -> None:
    assert failure(plain(tmp_path / "blank.png"), "Text (OCR)", tmp_path / "out") == "No text was found"


# ── document scan ────────────────────────────────────────────────────────


def test_scan_flattens_and_cleans_a_photographed_page(tmp_path: Path) -> None:
    source = page_photo(tmp_path / "page.jpg")
    color = Image.open(ok(source, "Scan to JPG", tmp_path / "out"))
    # Flattened to the page's own shape (850 × 1100), not the photo's.
    assert color.height / color.width == pytest.approx(1100 / 850, rel=0.06)
    pixels = np.asarray(color.convert("L"))
    assert np.median(pixels) > 240  # paper turned white
    assert pixels.min() < 80  # ink kept
    bw = Image.open(ok(source, "Scan to JPG B&W", tmp_path / "out"))
    assert np.asarray(bw.convert("RGB")).std(axis=2).max() < 12  # grey only


def test_scan_to_pdf_is_searchable_and_groups_into_one(tmp_path: Path) -> None:
    import pypdfium2 as pdfium

    source = page_photo(tmp_path / "page.jpg")
    with pdfium.PdfDocument(str(ok(source, "Scan to PDF", tmp_path / "out"))) as document:
        assert "Meeting notes" in document[0].get_textpage().get_text_range()
    second = page_photo(tmp_path / "page 2.jpg")
    output = ok(source, "One scanned PDF B&W", tmp_path / "out", members=[source, second])
    with pdfium.PdfDocument(str(output)) as document:
        assert len(document) == 2


def test_scan_reports_a_photo_without_a_document(tmp_path: Path) -> None:
    rng = np.random.default_rng(5)
    noise = tmp_path / "noise.jpg"
    Image.fromarray(rng.integers(0, 255, (900, 1200, 3), dtype=np.uint8)).save(noise)
    assert failure(noise, "Scan to JPG", tmp_path / "out") == "No document edges were found"


# ── codes ────────────────────────────────────────────────────────────────


def test_codes_to_text_and_one_csv(tmp_path: Path) -> None:
    first = qr_code(tmp_path / "ticket.png", "https://wemute.app/ticket/42")
    second = qr_code(tmp_path / "wifi.png", "WIFI:S:Wemute;T:WPA;P:secret;;")
    line = ok(first, "Read codes", tmp_path / "out").read_text(encoding="utf-8").strip()
    assert line == "ticket.png\tQR\thttps://wemute.app/ticket/42"
    output = ok(first, "Codes CSV", tmp_path / "out", members=[first, second])
    with output.open(encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == ["File", "Type", "Content"]
    assert ["wifi.png", "QR", "WIFI:S:Wemute;T:WPA;P:secret;;"] in rows


def test_codes_report_an_image_without_codes(tmp_path: Path) -> None:
    message = failure(plain(tmp_path / "plain.png"), "Read codes", tmp_path / "out")
    assert message == "No QR code or barcode was found"


# ── photo tools ──────────────────────────────────────────────────────────


def test_blur_faces_hides_every_face_and_nothing_else(tmp_path: Path) -> None:
    from wemute.engines.photo import detect_faces

    original = Image.open(NASA).convert("RGB")
    faces = detect_faces(original)
    assert len(faces) >= 6
    for tool in ("Blur faces", "Pixelate faces"):
        output = ok(NASA, tool, tmp_path / "out")
        assert output.name == f"{NASA.stem} faces hidden.jpg" or output.stem.startswith(NASA.stem)
        edited = Image.open(output).convert("RGB")
        assert edited.size == original.size
        # A mosaic may still read as a head shape, never as the person.
        assert len(detect_faces(edited)) <= (0 if tool == "Blur faces" else len(faces) // 3)
        a, b = np.asarray(original, dtype=float), np.asarray(edited, dtype=float)
        x0, y0, x1, y1 = (round(value) for value in faces[0])
        assert np.abs(a[y0:y1, x0:x1] - b[y0:y1, x0:x1]).mean() > 8
        assert np.abs(a[:60] - b[:60]).mean() < 2  # the wall above everyone is untouched


def test_faces_report_a_photo_without_people(tmp_path: Path) -> None:
    assert failure(tilted_sea(tmp_path / "sea.jpg"), "Blur faces", tmp_path / "out") == "No faces were found"


def test_id_photo_replaces_the_background(tmp_path: Path) -> None:
    edited = Image.open(ok(NASA, "ID photo blue", tmp_path / "out")).convert("RGB")
    pixels = np.asarray(edited, dtype=int)
    corner = pixels[:40, :40].reshape(-1, 3).mean(axis=0)
    assert np.abs(corner - (67, 142, 219)).max() < 12
    original = np.asarray(Image.open(NASA).convert("RGB"), dtype=int)
    middle = (slice(pixels.shape[0] // 2 - 20, pixels.shape[0] // 2 + 20), slice(pixels.shape[1] // 2 - 20, pixels.shape[1] // 2 + 20))
    assert np.abs(pixels[middle] - original[middle]).mean() < 10  # people stay as they were
    assert failure(plain(tmp_path / "wall.png"), "ID photo white", tmp_path / "out") == "No person was found"


def test_smart_crop_keeps_ratio_and_the_subject(tmp_path: Path) -> None:
    rng = np.random.default_rng(3)
    canvas = Image.fromarray(np.clip(rng.normal(120, 6, (1000, 1600, 3)), 0, 255).astype(np.uint8))
    ImageDraw.Draw(canvas).ellipse((1150, 380, 1450, 680), fill=(220, 30, 40))
    source = tmp_path / "ball.png"
    canvas.save(source)
    for tool, ratio in (("Smart crop 1:1", 1.0), ("Smart crop 9:16", 9 / 16)):
        cropped = Image.open(ok(source, tool, tmp_path / "out")).convert("RGB")
        assert cropped.width / cropped.height == pytest.approx(ratio, abs=0.01)
        assert cropped.height == 1000  # full resolution, never upscaled or shrunk
        red = np.asarray(cropped, dtype=int)
        assert ((red[..., 0] > 180) & (red[..., 1] < 80)).sum() > 0.9 * np.pi * 150**2


def test_straighten_levels_a_tilted_horizon(tmp_path: Path) -> None:
    from wemute.engines.photo import horizon_angle

    source = tilted_sea(tmp_path / "sea.jpg", 4.0)
    assert horizon_angle(Image.open(source)) == pytest.approx(4.0, abs=0.6)
    output = ok(source, "Straighten", tmp_path / "out")
    straightened = Image.open(output)
    assert straightened.width < Image.open(source).width  # cropped to the inscribed rectangle
    remaining = horizon_angle(straightened)
    assert remaining is None or abs(remaining) < 0.5
    message = failure(output, "Straighten", tmp_path / "again")
    assert "nothing to straighten" in message


def test_auto_enhance_lifts_a_dull_photo(tmp_path: Path) -> None:
    original = Image.open(NASA).convert("RGB")
    dull = Image.blend(original, Image.new("RGB", original.size, (128, 128, 128)), 0.45)
    source = tmp_path / "dull.jpg"
    dull.save(source, quality=92, icc_profile=Image.open(NASA).info.get("icc_profile"))
    output = ok(source, "Auto enhance", tmp_path / "out")
    enhanced = Image.open(output)
    assert enhanced.size == dull.size
    assert enhanced.info.get("icc_profile")  # colour profile kept
    assert np.asarray(enhanced.convert("L")).std() > np.asarray(dull.convert("L")).std() * 1.1


def test_outputs_never_overwrite(tmp_path: Path) -> None:
    out = tmp_path / "out"
    first = ok(NASA, "Auto enhance", out)
    second = ok(NASA, "Auto enhance", out)
    assert first != second and first.exists() and second.exists()


# ── blurry & duplicates ──────────────────────────────────────────────────


def test_triage_copies_duplicates_and_blurry_photos_into_a_report(tmp_path: Path) -> None:
    original = Image.open(NASA).convert("RGB")
    photos = tmp_path / "photos"
    photos.mkdir()
    best = photos / "astronauts.jpg"
    original.save(best, quality=92)
    smaller = photos / "astronauts small.jpg"
    original.resize((original.width * 3 // 4, original.height * 3 // 4)).save(smaller, quality=70)
    blurry = photos / "sea blurry.jpg"
    Image.open(tilted_sea(tmp_path / "sea.jpg")).filter(ImageFilter.GaussianBlur(9)).save(blurry)
    page = page_photo(photos / "page.jpg")
    sources = [best, smaller, blurry, page]
    before = {path: path.read_bytes() for path in sources}

    output, result = run_tool(best, "Blurry & duplicates", tmp_path / "out", members=sources)
    assert output is not None, result.stderr
    assert "wemute_note=Found 1 duplicates and 1 blurry photos" in result.stdout
    assert output.name == "Photo review" and output.is_dir()
    assert sorted(path.name for path in (output / "Duplicates").iterdir()) == ["astronauts small.jpg"]
    assert sorted(path.name for path in (output / "Blurry").iterdir()) == ["sea blurry.jpg"]
    with (output / "Report.csv").open(encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    flagged = {row["File"] for row in rows if row["Suggestion"] == "Consider deleting"}
    assert flagged == {"Duplicates/astronauts small.jpg", "Blurry/sea blurry.jpg"}
    assert {row["File"] for row in rows if row["Suggestion"] == "Keep"} == {"astronauts.jpg", "page.jpg"}
    assert all(path.read_bytes() == data for path, data in before.items())  # originals untouched


# ── Quick Look ───────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "make",
    [
        lambda d: fx.make_docx(d / "report.docx"),
        lambda d: fx.make_xlsx(d / "budget.xlsx"),
        lambda d: fx.make_pptx(d / "deck.pptx"),
        lambda d: fx.make_pdf(d / "paper.pdf"),
        lambda d: fx.make_svg(d / "logo.svg"),
        lambda d: fx.make_zip(d / "assets.zip"),
        lambda d: fx.make_markdown(d / "notes.md"),
    ],
)
def test_quick_look_thumbnails_for_every_file_kind(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, make) -> None:
    import wemute.probe as probe

    monkeypatch.setattr(probe, "THUMBNAIL_ROOT", tmp_path / "thumbs")
    source = make(tmp_path)
    kind = probe_file(source, create_preview=False).kind
    thumbnail = create_thumbnail(source, kind)
    assert thumbnail and Path(thumbnail).exists()
    with Image.open(thumbnail) as image:
        assert min(image.size) >= 120
        assert np.asarray(image.convert("L")).std() > 2  # a real picture, not a blank tile
    assert create_thumbnail(source, kind) == thumbnail  # cached
